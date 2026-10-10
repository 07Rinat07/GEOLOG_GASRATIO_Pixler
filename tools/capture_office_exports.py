"""Create synthetic RU/KK/EN Office exports for Windows desktop acceptance.

The runner validates bytes, numeric values and OpenXML structure; it cannot certify
that Microsoft Excel or Word opened/rendered the documents on an operator's machine.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook

from geoworkbench.calculations.interval_statistics import calculate_interval_statistics
from geoworkbench.data.interval_statistics_export import (
    export_interval_statistics_csv,
    export_interval_statistics_xlsx,
)
from geoworkbench.data.report_document_export import export_report_docx, export_report_html
from geoworkbench.data.selection_export import export_selection_excel, export_selection_text
from geoworkbench.domain.models import CurveData, CurveMetadata, Dataset, DatasetKind, DepthDomain
from geoworkbench.services.localization import AppLanguage, Localizer
from geoworkbench.services.report_definition import (
    ReportDefinition,
    ReportIntervalContext,
    ReportIntervalMode,
    ReportIntervalSelection,
    ReportProfile,
    resolve_report_definition,
)

BOM = b"\xef\xbb\xbf"
HEADINGS = {
    AppLanguage.RU: "Параметры отчёта",
    AppLanguage.KK: "Есеп параметрлері",
    AppLanguage.EN: "Report parameters",
}


def _dataset() -> Dataset:
    dataset = Dataset(
        "office-acceptance", "Скважина — Ұңғыма — Well",
        DatasetKind.GTI, DepthDomain.MD,
        np.array([100.0, 101.0, 102.0, 103.0]),
    )
    dataset.curves["c1"] = CurveData(
        CurveMetadata("c1", "C1", "C1", "%", "Methane", dataset.dataset_id),
        np.array([0.0, np.nan, 2.5, 3.0]),
    )
    dataset.curves["rop"] = CurveData(
        CurveMetadata("rop", "ROP", "ROP", "m/h", None, dataset.dataset_id),
        np.array([10.0, 20.0, 30.0, 40.0]),
    )
    return dataset


def _validate_openxml(path: Path) -> str:
    with ZipFile(path) as archive:
        assert archive.testzip() is None, path
        for member in archive.namelist():
            if member.endswith((".xml", ".rels")):
                ElementTree.fromstring(archive.read(member))
        if path.suffix == ".docx":
            body = archive.read("word/document.xml").decode("utf-8")
            return body
    return ""


def create_office_acceptance_bundle(output_dir: Path) -> dict[str, object]:
    """Generate 18 files and verify their actual content, not just their existence."""
    output_dir.mkdir(parents=True, exist_ok=True)
    dataset = _dataset()
    statistics = calculate_interval_statistics(dataset, 100.0, 103.0, ("C1", "ROP"))
    assert len(statistics) == 2
    files: list[dict[str, object]] = []

    for language in AppLanguage:
        folder = output_dir / language.value
        folder.mkdir(parents=True, exist_ok=True)
        targets = {
            "interval.csv": folder / "interval.csv",
            "interval.xlsx": folder / "interval.xlsx",
            "interval.docx": folder / "interval.docx",
            "interval.html": folder / "interval.html",
            "statistics.csv": folder / "statistics.csv",
            "statistics.xlsx": folder / "statistics.xlsx",
        }
        definition = ReportDefinition(
            f"office-acceptance:{language.value}", f"{dataset.name} selection",
            ReportProfile.COMBINED, dataset.dataset_id, dataset.active_index_id or "",
            ReportIntervalSelection(ReportIntervalMode.SELECTION),
            language=language.value,
            curve_ids=("c1", "rop"),
            channel_mnemonics=("C1", "ROP", "H2S"),
        )
        report = resolve_report_definition(
            dataset, definition,
            context=ReportIntervalContext(selection_range=(100.0, 103.0)),
            require_curves=True,
        )
        export_selection_text(
            dataset, targets["interval.csv"], ["c1", "rop"], 100.0, 103.0,
            delimiter=",", language=language, unavailable_mnemonics=("H2S",),
        )
        export_selection_excel(
            dataset, targets["interval.xlsx"], ["c1", "rop"], 100.0, 103.0,
            language=language, unavailable_mnemonics=("H2S",),
        )
        export_report_docx(dataset, targets["interval.docx"], report, language=language)
        export_report_html(dataset, targets["interval.html"], report, language=language)
        localizer = Localizer.create(language)
        stats_options = {
            "interval_label": "100–103 m",
            "dataset_name": dataset.name,
            "display_names": {"C1": "Метан", "ROP": "ROP"},
            "language": language,
        }
        export_interval_statistics_csv(
            targets["statistics.csv"], statistics, **stats_options,
        )
        export_interval_statistics_xlsx(
            targets["statistics.xlsx"], statistics, **stats_options,
        )

        interval_csv = targets["interval.csv"]
        assert interval_csv.read_bytes().startswith(BOM)
        with interval_csv.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        assert len(rows) == 5
        assert rows[1] == ["100", "0", "10", "#N/A"]
        assert rows[2] == ["101", "", "20", "#N/A"]
        assert rows[3] == ["102", "2.5", "30", "#N/A"]
        assert rows[4] == ["103", "3", "40", "#N/A"]
        assert "C1 [%]" in rows[0][1]

        _validate_openxml(targets["interval.xlsx"])
        with load_workbook(targets["interval.xlsx"], data_only=False) as book:
            sheet = book["Data"]
            assert sheet["B2"].value == 0
            assert sheet["B3"].value is None
            assert sheet["B4"].value == 2.5
            assert sheet["D2"].value == "#N/A"
            assert book["Metadata"]["B6"].value == language.value
            assert all(
                cell.data_type != "f"
                for ws in book.worksheets for row in ws for cell in row
            )

        document_xml = _validate_openxml(targets["interval.docx"])
        assert HEADINGS[language] in document_xml
        assert "#N/A" in document_xml
        html = targets["interval.html"].read_text(encoding="utf-8")
        assert f'<html lang="{language.value}">' in html
        assert f"<h2>{HEADINGS[language]}</h2>" in html
        assert "#N/A" in html

        statistics_csv = targets["statistics.csv"]
        assert statistics_csv.read_bytes().startswith(BOM)
        with statistics_csv.open(encoding="utf-8-sig", newline="") as stream:
            stats_rows = list(csv.reader(stream))
        assert stats_rows[3][0] == localizer.text("statistics.parameter")
        methane = next(row for row in stats_rows[4:] if row[1] == "C1")
        assert methane[4:8] == ["3", "1", "1", "75.0"]
        _validate_openxml(targets["statistics.xlsx"])
        with load_workbook(targets["statistics.xlsx"], data_only=False) as book:
            sheet = book.active
            assert sheet is not None
            assert sheet["A4"].value == localizer.text("statistics.parameter")
            assert sheet["H5"].value == 75
            assert sheet["I5"].value == 0
            assert all(cell.data_type != "f" for row in sheet for cell in row)

        for name, target in targets.items():
            payload = target.read_bytes()
            files.append({
                "language": language.value,
                "name": name,
                "path": str(target.relative_to(output_dir)).replace("\\", "/"),
                "size_bytes": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest(),
            })

    manifest: dict[str, object] = {
        "schema": 1,
        "source": "synthetic; contains no well/customer data",
        "status": "generated_and_structurally_validated",
        "microsoft_office_desktop_review": "not_performed_by_ci",
        "files": files,
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8",
    )
    (output_dir / "OPEN_IN_MICROSOFT_OFFICE.txt").write_text(
        "Проверка в Microsoft Excel и Word на Windows (вручную)\n"
        "1. Откройте interval.csv и statistics.csv двойным щелчком в Excel;\n"
        "   убедитесь, что русский/казахский текст читается без кракозябр.\n"
        "2. Откройте interval.xlsx и statistics.xlsx; проверьте числа,\n"
        "   пропущенные значения, нули, границы интервала и заголовки.\n"
        "3. Откройте interval.docx в Word; проверьте кириллицу, таблицы\n"
        "   и переносы на печати. Откройте interval.html в браузере.\n"
        "4. Повторите для папок ru, kk и en. Зафиксируйте версию Office.\n"
        "5. Статус manifest.json означает только автоматическую проверку\n"
        "   содержимого: Windows CI НЕ запускал Microsoft Office.\n",
        encoding="utf-8",
    )
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    options = parser.parse_args()
    manifest = create_office_acceptance_bundle(options.output_dir)
    print(f"Validated {len(manifest['files'])} synthetic exports: {options.output_dir}")


if __name__ == "__main__":
    main()
