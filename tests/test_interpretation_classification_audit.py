from __future__ import annotations

from dataclasses import replace
from html.parser import HTMLParser
import json
import zipfile

from defusedxml.ElementTree import fromstring
import numpy as np
from openpyxl import load_workbook
from PySide6.QtGui import QTextDocument
import pytest

from geoworkbench.data.hydrocarbon_interpretation_export import (
    export_hydrocarbon_interpretation_docx,
    export_hydrocarbon_interpretation_xlsx,
)
from geoworkbench.data.hydrocarbon_interpretation_export_docx_polished import (
    export_polished_hydrocarbon_interpretation_docx,
)
from geoworkbench.domain.depth_interval import DepthInterval
from geoworkbench.domain.models import Dataset, DatasetKind, DepthDomain
from geoworkbench.project.session import ProjectSession
from geoworkbench.services.hydrocarbon_interpretation import (
    build_hydrocarbon_interpretation_report,
    hydrocarbon_interpretation_html,
)
from geoworkbench.services.interpretation_classification_audit import (
    CLASSIFICATION_AUDIT_DOCX_PART,
    CLASSIFICATION_AUDIT_META,
    CLASSIFICATION_AUDIT_SCHEMA,
    CLASSIFICATION_AUDIT_SHEET,
)


def _report_and_dataset():
    depth = np.arange(1000.0, 1100.0)
    dataset = Dataset("audit-well", "Audit well", DatasetKind.GTI, DepthDomain.MD, depth)
    gas = np.ones(depth.size)
    gas[40:43] = (80.0, 120.0, 90.0)
    dataset.upsert_curve("C1_NORM", gas, unit="normalized gas units", provenance="calculation:test")
    session = ProjectSession()
    session.add_dataset(dataset, "Audit well")
    interval = DepthInterval(1030.0, 1060.0)
    report = build_hydrocarbon_interpretation_report(session, depth_interval=interval)
    assert report.candidates
    candidate = report.candidates[0]
    report = replace(
        report,
        candidates=(replace(candidate, fluid_hypothesis="opus_gasomer_gassy_oil"),),
        suppressed_candidates=(replace(candidate, fluid_hypothesis="heavy_or_residual_oil"),),
    )
    return report, dataset


class _AuditMetaParser(HTMLParser):
    payload = ""

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if tag == "meta" and attributes.get("name") == CLASSIFICATION_AUDIT_META:
            self.payload = attributes["content"]


def _assert_classification(payload, report):
    assert payload["schema"] == CLASSIFICATION_AUDIT_SCHEMA
    assert payload["dataset_id"] == report.dataset_id
    assert payload["analysis_depth_interval"] == [1030.0, 1060.0]
    assert [item["fluid_hypothesis"] for item in payload["candidates"]] == [
        "opus_gasomer_gassy_oil", "heavy_or_residual_oil",
    ]
    assert [item["status"] for item in payload["candidates"]] == ["active", "suppressed"]
    assert [item["phase"] for item in payload["candidates"]] == ["liquid", "liquid"]
    assert payload["candidates"][0]["evidence"] == list(report.candidates[0].evidence)


def test_html_preserves_raw_classification_in_nonvisible_metadata(qapp):
    report, _ = _report_and_dataset()
    html = hydrocarbon_interpretation_html(report)
    parser = _AuditMetaParser()
    parser.feed(html)
    _assert_classification(json.loads(parser.payload), report)
    document = QTextDocument()
    document.setHtml(html)
    assert "жидкая УВ-фаза" in document.toPlainText()
    assert "opus_gasomer_gassy_oil" not in document.toPlainText()
    assert "heavy_or_residual_oil" not in document.toPlainText()


def test_html_audit_escapes_unknown_hypothesis_without_creating_markup():
    report, _ = _report_and_dataset()
    raw = 'plugin_oil</head><img src="x" onerror="bad()">'
    report = replace(report, candidates=(replace(report.candidates[0], fluid_hypothesis=raw),))
    html = hydrocarbon_interpretation_html(report)
    parser = _AuditMetaParser()
    parser.feed(html)
    candidate = json.loads(parser.payload)["candidates"][0]
    assert candidate["fluid_hypothesis"] == raw
    assert candidate["phase"] == "indeterminate"
    assert '<img src="x"' not in html


@pytest.mark.parametrize("exporter", [
    export_hydrocarbon_interpretation_docx, export_polished_hydrocarbon_interpretation_docx,
])
def test_docx_audit_part_preserves_classification_without_visible_raw_codes(tmp_path, exporter):
    report, dataset = _report_and_dataset()
    path = exporter(report, tmp_path / "audit.docx", dataset=dataset)
    with zipfile.ZipFile(path) as package:
        audit = fromstring(package.read(CLASSIFICATION_AUDIT_DOCX_PART))
        _assert_classification(json.loads(audit.text), report)
        relationships = fromstring(package.read("word/_rels/document.xml.rels"))
        assert any(
            item.get("Target") == "../" + CLASSIFICATION_AUDIT_DOCX_PART
            and item.get("Type", "").endswith("/customXml")
            for item in relationships
        )
        body = package.read("word/document.xml").decode()
        assert "opus_gasomer_gassy_oil" not in body
        assert "heavy_or_residual_oil" not in body
        assert "жидкая УВ-фаза" in body


def test_xlsx_hidden_audit_preserves_long_unicode_evidence_and_literal_chunks(tmp_path):
    report, dataset = _report_and_dataset()
    evidence = ("=" * 40_000 + "🛢" * 20_000 + " қазақша / русский",)
    report = replace(report, candidates=(replace(report.candidates[0], evidence=evidence),))
    path = export_hydrocarbon_interpretation_xlsx(report, dataset, tmp_path / "audit.xlsx")
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook[CLASSIFICATION_AUDIT_SHEET]
        assert sheet.sheet_state == "hidden"
        cells = [row[2] for row in list(sheet.iter_rows())[1:]]
        assert len(cells) > 1
        assert any(cell.value.startswith("=") for cell in cells)
        assert all(cell.data_type == "s" for cell in cells)
        assert all(len(cell.value.encode("utf-16-le")) // 2 <= 32767 for cell in cells)
        payload = json.loads("".join(cell.value for cell in cells))
        _assert_classification(payload, report)
        assert payload["candidates"][0]["evidence"] == list(evidence)
        visible = str([list(sheet.values) for sheet in workbook if sheet.sheet_state == "visible"])
        assert "opus_gasomer_gassy_oil" not in visible
        assert "heavy_or_residual_oil" not in visible
    finally:
        workbook.close()
