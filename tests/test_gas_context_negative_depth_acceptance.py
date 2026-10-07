from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from pathlib import Path
from typing import Protocol
from unicodedata import normalize
import xml.etree.ElementTree as ET
import zipfile

import fitz
import numpy as np
from openpyxl import load_workbook
import pytest

from geoworkbench.data.hydrocarbon_interpretation_export import (
    export_hydrocarbon_interpretation_docx, export_hydrocarbon_interpretation_xlsx,
)
from geoworkbench.data.hydrocarbon_interpretation_export_docx_polished import (
    export_polished_hydrocarbon_interpretation_docx,
)
from geoworkbench.domain.depth_interval import DepthInterval, DepthIntervalError
from geoworkbench.domain.gas_context_events import GasContextEvent, GasContextEventType
from geoworkbench.domain.models import CurveData, Dataset, DepthDomain
from geoworkbench.printing.gas_context_track import context_segments
from geoworkbench.printing.hydrocarbon_interpretation_report import export_hydrocarbon_interpretation_pdf
from geoworkbench.project.controller import ProjectController
from geoworkbench.project.session import ProjectSession
from geoworkbench.services.gas_context_event_editor import GasContextEventEditorController
from geoworkbench.services.gas_context_report_labels import gas_context_type_text
from geoworkbench.services.hydrocarbon_interpretation import (
    build_hydrocarbon_interpretation_report, build_opus_interpretation_report,
    hydrocarbon_interpretation_html,
)
from geoworkbench.services.hydrocarbon_interpretation_legacy import HydrocarbonInterpretationReport
from geoworkbench.services.localization import AppLanguage
from geoworkbench.ui.gas_context_event_dialog import GasContextEventDialog
from test_hydrocarbon_interpretation import _session


class ReportBuilder(Protocol):
    def __call__(
        self, session: ProjectSession, *, threshold: float = 4.0,
        depth_interval: DepthInterval | None = None,
    ) -> HydrocarbonInterpretationReport: ...


def _negative_session(domain: DepthDomain = DepthDomain.TVDSS) -> ProjectSession:
    original = _session().current_dataset
    assert original is not None
    dataset = Dataset(original.dataset_id, "Signed depth", original.kind, domain, original.depth - 1100.0)
    dataset.curves = {
        key: CurveData(replace(curve.metadata), curve.values.copy())
        for key, curve in original.curves.items()
    }
    dataset.upsert_curve("TG", np.full(dataset.depth.shape, 4.0), unit="%", provenance="source")
    session = ProjectSession()
    session.add_dataset(dataset, "Signed depth well")
    session.dirty = False
    return session


@pytest.mark.parametrize("language", list(AppLanguage))
@pytest.mark.parametrize("builder", [build_hydrocarbon_interpretation_report, build_opus_interpretation_report])
@pytest.mark.parametrize("extension", [".json", ".geologpkg"])
@pytest.mark.parametrize("domain", [DepthDomain.MD, DepthDomain.TVD, DepthDomain.TVDSS])
def test_editor_signed_depth_reopens_with_same_context_in_every_export(
    qapp: object, tmp_path: Path, language: AppLanguage, builder: ReportBuilder, extension: str,
    domain: DepthDomain,
) -> None:
    session = _negative_session(domain)
    interval = DepthInterval(-70.0, -40.0)
    editor = GasContextEventEditorController(session)
    event = editor.add(event_type=GasContextEventType.CONNECTION_GAS, top_depth=-61.0,
                       bottom_depth=-57.0, reported_total_gas=4.25, reported_unit="%")
    repeated = editor.add(event_type=GasContextEventType.CONNECTION_GAS, top_depth=-50.0,
                          bottom_depth=-48.0)
    assert editor.commit()
    package = tmp_path / ("signed-depth" + extension)
    ProjectController(session=session).save_project(package)
    reopened = ProjectController().open_project(package)
    dataset, well = reopened.current_dataset, reopened.current_well
    assert dataset is not None and well is not None
    before = deepcopy(well.gas_context_events)
    depths = dataset.depth.copy()
    arrays = {key: curve.values.copy() for key, curve in dataset.curves.items()}
    report = builder(reopened, threshold=3.0, depth_interval=interval)
    before_candidates = deepcopy((report.candidates, report.suppressed_candidates, report.gas_context_audit))
    assert report.analysis_depth_interval == interval
    assert report.gas_context_events == (event, repeated)
    assert all(item.depth_domain is domain for item in report.gas_context_events)
    if builder is build_hydrocarbon_interpretation_report:
        assert report.candidates == ()
        assert len(report.suppressed_candidates) == 1
    audit = next(row for row in report.gas_context_audit if row.event_id == event.event_id)
    assert audit.measured_total_gas is not None and audit.measured_total_gas.mean == 4.0
    assert audit.manual_total_gas == 4.25 and audit.qc_delta_vs_measured_mean == 0.25
    assert audit.measured_components
    segments = context_segments(report.gas_context_events, interval.top_depth, interval.bottom_depth)
    assert [(row.top_depth, row.bottom_depth) for row in segments] == [(-61.0, -57.0), (-50.0, -48.0)]
    html = hydrocarbon_interpretation_html(report, language)
    assert "-61–-57" in html and "-50–-48" in html
    assert event.event_id in html and repeated.event_id in html
    assert gas_context_type_text(event.event_type, language) in html

    for export in (export_hydrocarbon_interpretation_docx, export_polished_hydrocarbon_interpretation_docx):
        target = tmp_path / (export.__name__ + ".docx")
        export(report, target, dataset=dataset, language=language)
        with zipfile.ZipFile(target) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
            text = " ".join(element.text or "" for element in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
        assert event.event_id in text and repeated.event_id in text
        assert "-61" in text and "-57" in text

    target = export_hydrocarbon_interpretation_xlsx(report, dataset, tmp_path / "context.xlsx", language=language)
    workbook = load_workbook(target)
    try:
        sheet = workbook[{AppLanguage.RU: "Газовый контекст", AppLanguage.KK: "Газ контексті", AppLanguage.EN: "Gas context"}[language]]
        assert [sheet.cell(2, column).value for column in (2, 3, 7, 11, 13, 17)] == [-61, -57, 4, 4.25, 0.25, event.event_id]
        assert sheet["B2"].data_type == sheet["C2"].data_type == "n"
    finally:
        workbook.close()
    target = tmp_path / "context.pdf"
    export_hydrocarbon_interpretation_pdf(report, target, dataset=dataset, language=language)
    with fitz.open(target) as document:
        text = "".join(normalize("NFKC", " ".join(page.get_text() for page in document)).split())
        assert event.event_id in text and repeated.event_id in text
        assert "-61" in text and "-57" in text
    assert well.gas_context_events == before
    assert (report.candidates, report.suppressed_candidates, report.gas_context_audit) == before_candidates
    np.testing.assert_array_equal(dataset.depth, depths)
    for key, values in arrays.items():
        np.testing.assert_array_equal(dataset.curves[key].values, values)
    assert not reopened.dirty


@pytest.mark.parametrize("builder", [build_hydrocarbon_interpretation_report, build_opus_interpretation_report])
@pytest.mark.parametrize("bounds", [(-60, -70), (-101, -50), (-70, 0), (float("nan"), -50), (-70, float("inf"))])
def test_invalid_signed_analysis_interval_fails_without_mutation(builder: ReportBuilder, bounds: tuple[float, float]) -> None:
    session = _negative_session()
    dataset = session.current_dataset
    assert dataset is not None
    before = {key: curve.values.copy() for key, curve in dataset.curves.items()}
    with pytest.raises(DepthIntervalError):
        builder(session, depth_interval=DepthInterval(*bounds))
    assert not session.dirty
    for key, values in before.items():
        np.testing.assert_array_equal(dataset.curves[key].values, values)


@pytest.mark.parametrize("builder", [build_hydrocarbon_interpretation_report, build_opus_interpretation_report])
def test_signed_time_axis_is_not_a_depth_analysis_axis(builder: ReportBuilder) -> None:
    session = _negative_session(DepthDomain.TIME)
    with pytest.raises(DepthIntervalError, match="MD/TVD/TVDSS"):
        builder(session, depth_interval=DepthInterval(-70, -40))
    assert not session.dirty


@pytest.mark.parametrize("builder", [build_hydrocarbon_interpretation_report, build_opus_interpretation_report])
def test_signed_context_on_other_or_ambiguous_legacy_axis_does_not_apply(builder: ReportBuilder) -> None:
    session = _negative_session()
    well = session.current_well
    dataset = session.current_dataset
    assert well is not None and dataset is not None
    well.datasets["other-axis"] = Dataset("other-axis", "Other axis", dataset.kind, DepthDomain.MD, np.arange(100.0))
    well.gas_context_events = [
        GasContextEvent("wrong-md", GasContextEventType.CONNECTION_GAS, -61, -57, depth_domain=DepthDomain.MD),
        GasContextEvent("legacy-unknown", GasContextEventType.CONNECTION_GAS, -61, -57),
    ]
    before = deepcopy(well.gas_context_events)
    report = builder(session, depth_interval=DepthInterval(-70, -40))
    assert report.gas_context_events == () and report.gas_context_audit == ()
    assert well.gas_context_events == before
    assert not session.dirty


@pytest.mark.parametrize("bounds", [(-60, -70), (float("nan"), -50), (-70, float("inf"))])
@pytest.mark.parametrize("operation", ["add", "update"])
def test_invalid_signed_event_keeps_editor_and_project_unchanged(bounds: tuple[float, float], operation: str) -> None:
    session = _negative_session()
    editor = GasContextEventEditorController(session)
    original = editor.add(event_type=GasContextEventType.CONNECTION_GAS, top_depth=-61, bottom_depth=-57)
    assert editor.commit()
    session.dirty = False
    before = editor.registry
    with pytest.raises(ValueError):
        if operation == "add":
            editor.add(event_type=GasContextEventType.SWAB_GAS, top_depth=bounds[0], bottom_depth=bounds[1])
        else:
            editor.update(original.event_id, event_type=GasContextEventType.SWAB_GAS, top_depth=bounds[0], bottom_depth=bounds[1])
    assert editor.registry == before and not editor.changed
    assert session.current_well is not None
    assert session.current_well.gas_context_events == [original]
    assert not session.dirty


@pytest.mark.parametrize("language", list(AppLanguage))
def test_dialog_saves_signed_tvdss_without_clamping(qapp: object, language: AppLanguage) -> None:
    session = _negative_session()
    editor = GasContextEventEditorController(session)
    dialog = GasContextEventDialog(editor, language=language)
    try:
        dialog.type_input.setCurrentIndex(dialog.type_input.findData(GasContextEventType.CONNECTION_GAS))
        dialog.top_input.setValue(-61.125)
        dialog.bottom_input.setValue(-57.875)
        dialog._add()
        assert not session.dirty
        dialog._save()
        well = session.current_well
        assert well is not None
        assert len(well.gas_context_events) == 1
        event = well.gas_context_events[0]
        assert event.depth_domain is DepthDomain.TVDSS
        assert (event.top_depth, event.bottom_depth) == (-61.125, -57.875)
        assert session.dirty
    finally:
        dialog.close()
