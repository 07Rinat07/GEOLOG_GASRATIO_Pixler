from copy import deepcopy
from dataclasses import replace
import html
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook

import numpy as np
import pytest

from test_report_document_export import _resolved_report
from geoworkbench.printing.hydrocarbon_report_i18n import hydrocarbon_report_labels
from geoworkbench.printing.report_document_control import compact_report_footer
from geoworkbench.printing.report_visual_system import modern_oilfield_report_profile
from geoworkbench.project.controller import ProjectController
from geoworkbench.project.dataset_export_controller import DatasetExportController
from geoworkbench.project.masterlog_template_controller import MasterlogTemplateController
from geoworkbench.project.session import ProjectSession
from geoworkbench.services.localization import AppLanguage
from geoworkbench.services.report_definition import ReportDefinitionError, resolve_report_definition, ReportIntervalContext


@pytest.mark.parametrize('language', list(AppLanguage))
@pytest.mark.parametrize('date', ['', '2026-09-30'])
@pytest.mark.parametrize('use_form', [False, True])
def test_generic_office_control_after_reopen(qapp, tmp_path, language, date, use_form):
    dataset, report = _resolved_report()
    dataset.headers['DATE'] = 'Acquisition date'
    session = ProjectSession()
    session.add_dataset(dataset, 'Well A')
    forms = MasterlogTemplateController(session)
    template = forms.create('Controlled form')
    fields = {'header.document_number': 'DOC<&42', 'header.revision': '07',
              'header.status': 'Approved', 'header.report_date': date,
              'header.prepared_by': 'Engineer A', 'header.checked_by': 'Engineer B',
              'header.approved_by': 'Engineer C', 'header.confidentiality': 'Internal ' + 'Long text<& ' * 80,
              'header.interval': 'stale interval'}
    forms.update_header_fields(template.template_id, fields)
    unrelated = forms.create('Unrelated form')
    forms.update_header_fields(unrelated.template_id, {'header.document_number': 'UNRELATED'})
    definition = replace(report.definition, language=language.value,
                         form_kind='masterlog-template' if use_form else None,
                         form_id=template.template_id if use_form else None,
                         form_revision=f'version:{template.version}' if use_form else None)
    package = tmp_path / 'control.geologpkg'
    ProjectController(session=session).save_project(package)
    restored = ProjectController().open_project(package)
    dataset = restored.current_dataset
    report = resolve_report_definition(dataset, definition, context=ReportIntervalContext(selection_range=(100, 102)), require_curves=True)
    before = {key: curve.values.copy() for key, curve in dataset.curves.items()}
    saved_fields = deepcopy(restored.project.masterlog_templates[template.template_id].properties)
    controller = DatasetExportController(restored)
    html_target, docx_target = tmp_path / 'report.html', tmp_path / 'report.docx'
    controller.export_resolved_report_html(html_target, report)
    controller.export_resolved_report_docx(docx_target, report)
    xlsx_target = tmp_path / 'report.xlsx'
    controller.export_resolved_report_excel(xlsx_target, report, language=language)
    with zipfile.ZipFile(xlsx_target) as archive:
        assert archive.testzip() is None
    workbook = load_workbook(xlsx_target)
    control_sheet = workbook[{AppLanguage.RU: 'Реквизиты', AppLanguage.KK: 'Деректемелер', AppLanguage.EN: 'Document control'}[language]]
    control_text = ' '.join(str(cell.value) for row in control_sheet for cell in row if cell.value is not None)
    assert 'Well A' in control_text and '100.0 — 102.0 m' in control_text
    assert 'UNRELATED' not in control_text and 'stale interval' not in control_text
    assert (fields['header.document_number'] in control_text) == use_form
    assert (hydrocarbon_report_labels(language).report_date in control_text) == bool(use_form and date)
    assert control_sheet.page_setup.fitToWidth == 1 and control_sheet.page_setup.fitToHeight == 0
    assert control_sheet.print_title_rows == '$1:$3'
    assert control_sheet.oddFooter.left.text == modern_oilfield_report_profile().brand_wordmark.replace('&', '&&')
    data_sheet = workbook['Data']
    assert data_sheet['B2'].value == 0 and data_sheet['B2'].data_type == 'n'
    assert data_sheet['B3'].value is None
    assert data_sheet['C2'].value == '#N/A'
    assert data_sheet['B4'].value == 25 and data_sheet['B4'].data_type == 'n'
    assert data_sheet.print_title_rows == '$1:$1'
    assert data_sheet.print_title_cols == '$A:$A'
    assert data_sheet.page_setup.orientation == 'landscape'
    assert data_sheet.page_setup.fitToWidth == 0 and data_sheet.page_setup.scale == 100
    assert workbook['Parameters'].print_title_cols == '$A:$B'
    assert data_sheet['B2'].alignment.horizontal == 'right'
    workbook.close()
    markup = html_target.read_text()
    plain_html = html.unescape(markup)
    namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(docx_target) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith('.xml') or name.endswith('.rels'):
                ET.fromstring(archive.read(name))
        document = ET.fromstring(archive.read('word/document.xml'))
        footer = ET.fromstring(archive.read('word/footer.xml'))
        word_text = ' '.join(node.text or '' for node in document.findall('.//w:t', namespace))
        footer_text = ' '.join(node.text or '' for node in footer.findall('.//w:t', namespace))
        relationships = archive.read('word/_rels/document.xml.rels').decode()
        assert 'Target="footer.xml"' in relationships
        assert document.find('.//w:footerReference', namespace) is not None
        assert {node.get('{'+namespace['w']+'}instr') for node in footer.findall('.//w:fldSimple', namespace)} == {'PAGE', 'NUMPAGES'}
        assert document.findall('.//w:tblHeader', namespace)
        assert modern_oilfield_report_profile().palette.table_header.lstrip('#') in archive.read('word/document.xml').decode()
    for text in (plain_html, word_text):
        assert 'Well A' in text
        assert '100.0 — 102.0 m' in text
        assert 'UNRELATED' not in text and 'stale interval' not in text and 'Acquisition date' not in text
        for key in ('header.document_number', 'header.prepared_by', 'header.checked_by', 'header.approved_by', 'header.confidentiality'):
            assert (fields[key].strip() in text) == use_form
        assert (hydrocarbon_report_labels(language).report_date in text) == bool(use_form and date)
        assert (date in text) if date and use_form else True
    assert len(footer_text) < 260
    if use_form:
        assert '…' in footer_text
    assert 'DOC&lt;&amp;42' in markup if use_form else True
    assert modern_oilfield_report_profile().brand_wordmark in word_text
    assert modern_oilfield_report_profile().brand_wordmark in footer_text
    assert report.definition.content_sha256 in word_text
    assert '#N/A' in word_text and '—' in word_text
    assert saved_fields == restored.project.masterlog_templates[template.template_id].properties
    for key, values in before.items():
        np.testing.assert_array_equal(values, dataset.curves[key].values)


@pytest.mark.parametrize('reason', ['missing', 'changed'])
@pytest.mark.parametrize('format', ['html', 'docx', 'excel'])
def test_generic_export_rejects_changed_form_revision(tmp_path, reason, format):
    dataset, original = _resolved_report()
    session = ProjectSession()
    session.add_dataset(dataset, 'Well')
    forms = MasterlogTemplateController(session)
    template = forms.create('Form')
    definition = replace(original.definition, form_kind='masterlog-template', form_id=template.template_id,
                         form_revision=f'version:{template.version}')
    report = resolve_report_definition(dataset, definition, context=ReportIntervalContext(selection_range=(100, 102)))
    if reason == 'missing':
        forms.delete(template.template_id)
    else:
        forms.update_header_fields(template.template_id, {'header.revision': '08'})
    target = tmp_path / f'unchanged.{"xlsx" if format == "excel" else format}'
    target.write_bytes(b'original output')
    export = getattr(DatasetExportController(session), f'export_resolved_report_{format}')
    with pytest.raises(ReportDefinitionError, match='Ревизия'):
        export(target, report, overwrite=True)
    assert target.read_bytes() == b'original output'


def test_compact_footer_normalizes_and_bounds_each_value():
    from geoworkbench.printing.report_document_control import ReportDocumentControl
    control = ReportDocumentControl('', '', (), (), (), (), ('  first\nsecond  ', 'x' * 1000))
    assert compact_report_footer(control) == 'first second · ' + 'x' * 47 + '…'
    assert compact_report_footer(None) == ''


@pytest.mark.parametrize('language', list(AppLanguage))
def test_generic_xlsx_control_formula_like_values_are_literal(tmp_path, language):
    from geoworkbench.printing.hydrocarbon_interpretation_report_identity import InterpretationReportIdentity
    from geoworkbench.printing.report_document_control import report_document_control
    from geoworkbench.data.selection_export import export_selection_excel
    dataset, report = _resolved_report()
    identity = InterpretationReportIdentity('=TITLE', '', '@PROJECT', '+WELL', revision='-REV',
                                          document_number='=HYPERLINK("evil")', confidentiality='@PRIVATE')
    snapshot = report_document_control(identity, language)
    target = tmp_path / 'literal.xlsx'
    export_selection_excel(dataset, target, ['c1'], 100, 102, row_indices=report.interval.indices,
                           document_control=snapshot, language=language)
    workbook = load_workbook(target)
    sheet = workbook[{AppLanguage.RU: 'Реквизиты', AppLanguage.KK: 'Деректемелер', AppLanguage.EN: 'Document control'}[language]]
    values = [cell for row in sheet for cell in row if cell.value is not None]
    for value in ('=TITLE', '@PROJECT', '+WELL', '-REV', '=HYPERLINK("evil")', '@PRIVATE'):
        cell = next(cell for cell in values if str(cell.value).endswith(value))
        assert cell.data_type == 's'
    assert workbook['Data']['B2'].value == 0 and workbook['Data']['B2'].data_type == 'n'
    workbook.close()
