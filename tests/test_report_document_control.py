from dataclasses import FrozenInstanceError, replace
from contextlib import nullcontext
import zipfile
from xml.etree import ElementTree

import fitz
import numpy as np
import pytest
from openpyxl import load_workbook

from geoworkbench.data.hydrocarbon_interpretation_export_docx_polished import export_polished_hydrocarbon_interpretation_docx
from geoworkbench.data.hydrocarbon_interpretation_export_readable import export_readable_hydrocarbon_interpretation_xlsx
from geoworkbench.domain.depth_interval import DepthInterval
from geoworkbench.domain.report_composition import InterpretationReportComposition, with_report_header_fields
from geoworkbench.printing.hydrocarbon_report_i18n import hydrocarbon_report_labels
from geoworkbench.printing.hydrocarbon_interpretation_report import export_hydrocarbon_interpretation_pdf
from geoworkbench.printing.hydrocarbon_interpretation_report_identity import report_header_fields_from_identity
from geoworkbench.printing.report_document_control import report_document_control, resolved_report_identity
from geoworkbench.project.controller import ProjectController
from geoworkbench.project.interpretation_calculation_controller import InterpretationCalculationController
from geoworkbench.services.hydrocarbon_interpretation import build_hydrocarbon_interpretation_report
from geoworkbench.services.localization import AppLanguage
from geoworkbench.ui.interpretation_report_workspace_final import InterpretationReportWorkspace
from test_interpretation_report_charts import _session_with_report_curves
from test_interpretation_report_identity import _manual_identity


SHEETS = {AppLanguage.RU: 'Реквизиты', AppLanguage.KK: 'Деректемелер', AppLanguage.EN: 'Document control'}


@pytest.mark.parametrize('language', list(AppLanguage))
@pytest.mark.parametrize('profile', ['standard', 'opus'])
@pytest.mark.parametrize('date', ['', '06.10.2026'])
def test_same_control_snapshot_reaches_pdf_docx_and_xlsx(qapp, tmp_path, language, profile, date):
    session = _session_with_report_curves(depth_span=30, samples=61)
    dataset = session.current_dataset
    originals = {key: curve.values.copy() for key, curve in dataset.curves.items()}
    report = replace(build_hydrocarbon_interpretation_report(session), report_profile=profile,
                     analysis_depth_interval=DepthInterval(1305, 1320))
    generated = report.generated_at
    identity = replace(_manual_identity(), report_title='Client control', project_name='Client project',
                       well_name='Client well', document_number='DOC-017', revision='07',
                       document_status='Approved', prepared_by='Engineer A', checked_by='Engineer B',
                       approved_by='Engineer C', report_date=date, interval='stale interval')
    resolved = resolved_report_identity(report, identity, language)
    snapshot = report_document_control(resolved, language)
    assert resolved.interval == report.analysis_depth_interval.formatted(report.depth_unit)
    pdf = export_hydrocarbon_interpretation_pdf(report, tmp_path / 'control.pdf', dataset=dataset,
                                              language=language, identity=identity)
    with fitz.open(pdf) as document:
        pdf_text = document[0].get_text()
    docx = export_polished_hydrocarbon_interpretation_docx(report, tmp_path / 'control.docx',
                                                        dataset=dataset, identity=identity, language=language)
    with zipfile.ZipFile(docx) as package:
        xml = ElementTree.fromstring(package.read('word/document.xml'))
    word_text = '\n'.join(node.text or '' for node in xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
    xlsx = export_readable_hydrocarbon_interpretation_xlsx(report, dataset, tmp_path / 'control.xlsx',
                                                        language=language, identity=identity)
    workbook = load_workbook(xlsx)
    sheet = workbook[SHEETS[language]]
    # The visible report and the document-control sheet must use the same
    # edited passport values, even if raw report.project_name is stale.
    interpretation = workbook[hydrocarbon_report_labels(language).sheet_interpretation]
    assert interpretation['A1'].value == identity.report_title
    assert interpretation['B2'].value == 'Client project'
    assert interpretation['F2'].value == 'Client well'
    assert interpretation['B3'].value == identity.dataset_name
    assert interpretation['F3'].value == resolved.interval
    assert interpretation['B2'].value != report.project_name
    excel_rows = [(row[0], row[1]) for row in sheet.iter_rows(min_row=4, values_only=True) if row[1] is not None]
    assert excel_rows == list(snapshot.available_rows)
    assert sheet["A3"].value == identity.report_subtitle
    for value in ['DOC-017', '07', 'Approved', 'Engineer A', 'Engineer B', 'Engineer C', resolved.interval]:
        assert value in pdf_text
        assert value in word_text
    for text in (pdf_text, word_text, str(list(sheet.values))):
        assert generated not in text
        assert 'stale interval' not in text
    if date:
        assert date in pdf_text and date in word_text
    else:
        assert len(snapshot.control) == 3
    assert report.generated_at == generated
    for key, values in originals.items():
        np.testing.assert_array_equal(dataset.curves[key].values, values)
    workbook.close()


@pytest.mark.parametrize('language', list(AppLanguage))
@pytest.mark.parametrize('stored_profile', ['standard', 'opus'])
def test_workspace_xlsx_uses_saved_reopened_language_header(qapp, tmp_path, monkeypatch, language, stored_profile):
    session = _session_with_report_curves(depth_span=30, samples=61)
    dataset = session.current_dataset
    identity = replace(_manual_identity(), document_number='SAVED-42', report_date='')
    composition = with_report_header_fields(InterpretationReportComposition(), language.value,
                                           report_header_fields_from_identity(identity, stored_profile))
    session.report_compositions[dataset.dataset_id] = composition
    project = tmp_path / 'saved.geologpkg'
    ProjectController(session=session).save_project(project)
    restored = ProjectController().open_project(project)
    before = dict(restored.report_compositions)
    workspace = InterpretationReportWorkspace(InterpretationCalculationController(restored), language=language)
    report = build_hydrocarbon_interpretation_report(restored)
    target = tmp_path / 'workspace.xlsx'
    monkeypatch.setattr(workspace, '_require_report', lambda: report)
    monkeypatch.setattr(workspace, '_choose_target', lambda *args: target)
    monkeypatch.setattr(workspace, '_report_export_progress', lambda *args: nullcontext())
    monkeypatch.setattr(workspace, '_update_report_export_progress', lambda *args: None)
    monkeypatch.setattr(workspace, '_show_export_success', lambda *args: None)
    monkeypatch.setattr(workspace, '_show_export_error', lambda error: pytest.fail(str(error)))
    try:
        workspace._export_xlsx()
        workbook = load_workbook(target)
        values = str(list(workbook[SHEETS[language]].values))
        assert ('SAVED-42' in values) == (stored_profile == 'standard')
        assert workspace._report_interval(report) in values
        if stored_profile == 'standard':
            assert workbook[SHEETS[language]]["A3"].value == identity.report_subtitle
        assert restored.report_compositions == before
        workbook.close()
    finally:
        workspace.close()


def test_document_control_immutable_and_excel_formula_like_text_is_literal(tmp_path):
    session = _session_with_report_curves(depth_span=30, samples=61)
    report = build_hydrocarbon_interpretation_report(session)
    identity = replace(_manual_identity(), document_number='=2+2', report_date='  ', remarks='@unsafe')
    snapshot = report_document_control(identity, AppLanguage.EN)
    with pytest.raises(FrozenInstanceError):
        snapshot.title = 'changed'
    assert len(snapshot.control) == 3
    target = export_readable_hydrocarbon_interpretation_xlsx(report, session.current_dataset,
        tmp_path / 'literal.xlsx', language=AppLanguage.EN, identity=identity)
    workbook = load_workbook(target)
    sheet = workbook['Document control']
    assert not any(cell.data_type == 'f' for row in sheet for cell in row)
    assert any('=2+2' in str(cell.value) for row in sheet for cell in row)
    assert not any(cell.data_type == 'f' for row in workbook['HC interpretation'] for cell in row)
    workbook.close()


def test_xlsx_edited_title_is_never_executed_as_excel_formula(tmp_path) -> None:
    session = _session_with_report_curves(depth_span=30, samples=61)
    report = build_hydrocarbon_interpretation_report(session)
    edited = replace(
        _manual_identity(), report_title='=HYPERLINK("https://example.org", "click")',
        project_name='+777', well_name='@M-1', dataset_name='-123',
    )
    path = export_readable_hydrocarbon_interpretation_xlsx(
        report, session.current_dataset, tmp_path / 'safe-title.xlsx',
        language=AppLanguage.EN, identity=edited,
    )
    workbook = load_workbook(path)
    try:
        sheet = workbook['HC interpretation']
        assert all(sheet[cell].data_type != 'f' for cell in ('A1', 'B2', 'F2', 'B3'))
        for cell in ('A1', 'B2', 'F2', 'B3'):
            assert str(sheet[cell].value).lstrip("'") == {
                'A1': edited.report_title, 'B2': edited.project_name,
                'F2': edited.well_name, 'B3': edited.dataset_name,
            }[cell]
    finally:
        workbook.close()

