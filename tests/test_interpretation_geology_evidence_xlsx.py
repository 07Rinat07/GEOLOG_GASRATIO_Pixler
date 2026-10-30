"""Opt-in measured geology in both public Excel entry points."""
from dataclasses import replace

import numpy as np
from openpyxl import load_workbook
import pytest

from geoworkbench.data.hydrocarbon_interpretation_export import (
    HydrocarbonInterpretationExportError, export_hydrocarbon_interpretation_xlsx,
)
from geoworkbench.data import hydrocarbon_interpretation_export_readable as readable
from geoworkbench.services.interpretation_geology_interval_evidence_html import geology_interval_evidence_table
from geoworkbench.services.interpretation_geology_interval_evidence import (
    correlate_gas_interval_with_well_geology,
)
from geoworkbench.services.interpretation_geology_field_handoff import (
    confirm_geology_field_selection,
)
from geoworkbench.services.localization import AppLanguage
from test_interpretation_geology_evidence_signed_pdf import _actual_source


@pytest.mark.parametrize("exporter", (export_hydrocarbon_interpretation_xlsx, readable.export_readable_hydrocarbon_interpretation_xlsx))
@pytest.mark.parametrize("language", tuple(AppLanguage))
def test_selected_excel_matches_shared_projection(tmp_path, language, exporter):
    session, report, evidence = _actual_source()
    dataset = session.current_dataset
    depth = dataset.depth.copy()
    curves = {key: curve.values.copy() for key, curve in dataset.curves.items()}
    plain = exporter(report, dataset, tmp_path / "plain.xlsx", language=language)
    selected = exporter(report, dataset, tmp_path / "selected.xlsx", language=language, geology_evidence=(evidence,))
    name = readable._GEOLOGY_EVIDENCE_SHEET_NAMES[language]
    with_book = load_workbook(selected)
    without_book = load_workbook(plain)
    try:
        assert name not in without_book.sheetnames
        assert with_book.sheetnames == without_book.sheetnames[:2] + [name] + without_book.sheetnames[2:]
        rows = list(with_book[name].values)
        projection = geology_interval_evidence_table(evidence, language)
        assert rows[0][0] == projection.title
        assert rows[1][0] == projection.intro
        for row in (*projection.coverage_rows, projection.headings, *projection.rows):
            assert tuple(row) in [tuple(value for value in actual[:len(row)]) for actual in rows]
        for sheet_name in without_book.sheetnames:
            assert list(with_book[sheet_name].values) == list(without_book[sheet_name].values)
    finally:
        with_book.close()
        without_book.close()
    np.testing.assert_array_equal(dataset.depth, depth)
    for key, values in curves.items():
        np.testing.assert_array_equal(dataset.curves[key].values, values)


@pytest.mark.parametrize("change", ("dataset", "axis", "candidate", "duplicate", "active_axis"))
def test_invalid_selection_preserves_destination(tmp_path, change):
    session, report, evidence = _actual_source()
    dataset = session.current_dataset
    selected = (evidence,)
    if change == "dataset":
        selected = (replace(evidence, dataset_id="foreign"),)
    elif change == "axis":
        selected = (replace(evidence, depth_axis_id="foreign"),)
    elif change == "candidate":
        selected = (replace(evidence, gas_top=1003.0),)
    elif change == "duplicate":
        selected = (evidence, evidence)
    else:
        report = replace(report, depth_axis_id="foreign")
        selected = (replace(evidence, depth_axis_id="foreign"),)
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"existing workbook")
    with pytest.raises(ValueError):
        export_hydrocarbon_interpretation_xlsx(report, dataset, target, overwrite=True, geology_evidence=selected)
    assert target.read_bytes() == b"existing workbook"


@pytest.mark.parametrize("source", ("=1+1", "+SUM(A1)", "-1+2", "@SUM(A1)", " \t=1+1"))
def test_source_ids_are_text_not_formulas(tmp_path, source):
    session, report, evidence = _actual_source()
    evidence = replace(evidence, intersections=(replace(evidence.intersections[0], source_id=source),))
    path = export_hydrocarbon_interpretation_xlsx(report, session.current_dataset, tmp_path / "safe.xlsx", language=AppLanguage.EN, geology_evidence=(evidence,))
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["Gas and geology MD"]
        cells = [cell for row in sheet for cell in row]
        assert any(cell.value == "'" + source for cell in cells)
        assert all(cell.data_type != "f" for cell in cells)
    finally:
        workbook.close()


def test_row_limit_preserves_destination(tmp_path, monkeypatch):
    session, report, evidence = _actual_source()
    projection = geology_interval_evidence_table(evidence, AppLanguage.EN)
    rows = 7 + len(projection.coverage_rows) + len(projection.rows)
    monkeypatch.setattr(readable, "_EXCEL_MAX_ROWS", rows - 1)
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"existing workbook")
    with pytest.raises(HydrocarbonInterpretationExportError, match="Geology evidence exceeds"):
        export_hydrocarbon_interpretation_xlsx(report, session.current_dataset, target, language=AppLanguage.EN, overwrite=True, geology_evidence=(evidence,))
    assert target.read_bytes() == b"existing workbook"


@pytest.mark.parametrize("exporter", (export_hydrocarbon_interpretation_xlsx, readable.export_readable_hydrocarbon_interpretation_xlsx))
@pytest.mark.parametrize("field", ("description", "protected_id"))
@pytest.mark.parametrize("length", (32767, 32768))
@pytest.mark.parametrize("language", tuple(AppLanguage))
def test_cell_text_limit_checked_before_target_preparation(tmp_path, exporter, field, length, language):
    session, report, evidence = _actual_source()
    observation = evidence.intersections[0]
    if field == "description":
        label = {AppLanguage.RU: "Описание", AppLanguage.KK: "Сипаттама", AppLanguage.EN: "Description"}[language] + ": "
        value = label + "x" * (length - len(label))
        observation = replace(observation, attributes=(("description", value[len(label):]),))
    else:
        # The stored apostrophe counts toward the limit as well.
        value = "'=" + "x" * (length - 2)
        observation = replace(observation, source_id=value[1:])
    evidence = replace(evidence, intersections=(observation,))
    target = tmp_path / "new-parent" / "observed.xlsx"
    if length > 32767:
        with pytest.raises(HydrocarbonInterpretationExportError, match="cell text limit"):
            exporter(report, session.current_dataset, target, language=language, geology_evidence=(evidence,))
        assert not target.parent.exists()
        target = tmp_path / "existing.xlsx"
        target.write_bytes(b"existing workbook")
        with pytest.raises(HydrocarbonInterpretationExportError, match="cell text limit"):
            exporter(report, session.current_dataset, target, overwrite=True, language=language, geology_evidence=(evidence,))
        assert target.read_bytes() == b"existing workbook"
    else:
        exporter(report, session.current_dataset, target, language=language, geology_evidence=(evidence,))
        workbook = load_workbook(target)
        try:
            assert value in [cell.value for row in workbook[readable._GEOLOGY_EVIDENCE_SHEET_NAMES[language]] for cell in row]
        finally:
            workbook.close()


def test_multi_interval_geology_has_no_global_filter(tmp_path):
    from test_interpretation_geology_evidence_operator_ui import _candidate
    from geoworkbench.services.interpretation_geology_interval_evidence import correlate_gas_interval_with_well_geology
    session, report, evidence = _actual_source()
    second = correlate_gas_interval_with_well_geology(session, top_depth=1001.0, bottom_depth=1010.0)
    report = replace(report, candidates=(*report.candidates, _candidate(1001.0, 1010.0)))
    target = export_hydrocarbon_interpretation_xlsx(
        report, session.current_dataset, tmp_path / "multiple.xlsx", language=AppLanguage.EN,
        geology_evidence=(evidence, second),
    )
    workbook = load_workbook(target)
    try:
        sheet = workbook["Gas and geology MD"]
        assert sheet.auto_filter.ref is None
        assert not sheet.tables
        cells = [cell.value for row in sheet for cell in row]
        for selected in (evidence, second):
            assert geology_interval_evidence_table(selected, AppLanguage.EN).interval in cells
        # Ordinary single-table auxiliary sheets retain their filter.
        assert workbook.worksheets[1].auto_filter.ref is not None
    finally:
        workbook.close()



def _approved_excel_fields(session, report):
    evidence = correlate_gas_interval_with_well_geology(
        session, top_depth=1002.0, bottom_depth=1010.0,
    )
    key = ("approved-excel", report.dataset_id, report.depth_axis_id)
    selected = confirm_geology_field_selection(
        session, report, (evidence,),
        (("calcimetry", "calcite_percent"), ("calcimetry", "dolomite_percent")),
        source_key=key,
    )
    return selected, key


@pytest.mark.parametrize(
    "exporter",
    (export_hydrocarbon_interpretation_xlsx, readable.export_readable_hydrocarbon_interpretation_xlsx),
)
@pytest.mark.parametrize("language", tuple(AppLanguage))
def test_confirmed_excel_fields_have_exact_md_coverage_and_legacy_unchanged(
    tmp_path, exporter, language,
):
    session, report, _ = _actual_source()
    dataset = session.current_dataset
    assert dataset is not None
    session.current_well.cuttings[0].calcite_percent = 0.0
    session.current_well.cuttings[0].dolomite_percent = None
    approved, key = _approved_excel_fields(session, report)
    before = repr(session.current_well)
    plain = exporter(report, dataset, tmp_path / "plain.xlsx", language=language)
    selected = exporter(
        report, dataset, tmp_path / "selected.xlsx", language=language,
        session=session, geology_field_selection=approved,
        geology_source_key=key,
    )
    projection = geology_interval_evidence_table(
        approved.evidence[0], language, fields=approved.fields,
    )
    with_book = load_workbook(selected)
    without_book = load_workbook(plain)
    try:
        name = readable._GEOLOGY_EVIDENCE_SHEET_NAMES[language]
        assert name not in without_book.sheetnames
        rows = list(with_book[name].values)
        assert projection.field_coverage_title in [row[0] for row in rows]
        assert projection.field_coverage_headings in rows
        for row in projection.field_coverage_rows:
            assert row in [tuple(actual[:len(row)]) for actual in rows]
        assert "calcimetry" not in [str(cell) for r in without_book.worksheets[0] for cell in r]
        for sheet_name in without_book.sheetnames:
            assert list(without_book[sheet_name].values) == list(with_book[sheet_name].values)
    finally:
        with_book.close()
        without_book.close()
    assert repr(session.current_well) == before


@pytest.mark.parametrize(
    "exporter",
    (export_hydrocarbon_interpretation_xlsx, readable.export_readable_hydrocarbon_interpretation_xlsx),
)
@pytest.mark.parametrize(
    "change",
    ("missing_session", "missing_key", "wrong_key", "stale_geology",
     "foreign_report", "invalid_field", "mismatched_evidence", "unconfirmed_key"),
)
def test_excel_field_source_refusal_preserves_existing_workbook(
    tmp_path, exporter, change,
):
    session, report, evidence = _actual_source()
    approved, key = _approved_excel_fields(session, report)
    kwargs = {
        "session": session, "geology_field_selection": approved,
        "geology_source_key": key,
    }
    if change == "missing_session":
        kwargs.pop("session")
    elif change == "missing_key":
        kwargs.pop("geology_source_key")
    elif change == "wrong_key":
        kwargs["geology_source_key"] = ("other-source",)
    elif change == "stale_geology":
        session.current_well.cuttings[0].description = "changed after selection"
    elif change == "foreign_report":
        report = replace(report)
    elif change == "invalid_field":
        kwargs["geology_field_selection"] = replace(
            approved, fields=(("calcimetry", "unknown_assay"),),
        )
    elif change == "mismatched_evidence":
        kwargs["geology_evidence"] = (replace(evidence, well_id="foreign"),)
    else:
        kwargs = {"geology_source_key": key}
    target = tmp_path / "approved.xlsx"
    target.write_bytes(b"previous approved workbook")
    with pytest.raises(ValueError):
        exporter(report, session.current_dataset, target, overwrite=True, **kwargs)
    assert target.read_bytes() == b"previous approved workbook"


@pytest.mark.parametrize(
    "exporter",
    (export_hydrocarbon_interpretation_xlsx, readable.export_readable_hydrocarbon_interpretation_xlsx),
)
def test_excel_field_approval_revoked_during_progress_never_replaces_workbook(
    tmp_path, exporter,
):
    session, report, _ = _actual_source()
    approved, key = _approved_excel_fields(session, report)
    target = tmp_path / "approved.xlsx"
    target.write_bytes(b"original protected workbook")

    def edit_on_progress(_title, progress, _total):
        if progress == 95:
            session.current_well.cuttings[0].description = "changed during save"

    with pytest.raises(ValueError):
        exporter(
            report, session.current_dataset, target, overwrite=True,
            session=session, geology_field_selection=approved,
            geology_source_key=key, progress=edit_on_progress,
        )
    assert target.read_bytes() == b"original protected workbook"
