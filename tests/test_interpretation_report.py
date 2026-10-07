from dataclasses import replace
import zipfile

import fitz
import numpy as np
import pytest
from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QPalette
from PySide6.QtWidgets import QPushButton, QTextBrowser, QTextEdit

from geoworkbench.domain.models import Dataset, DatasetKind, DepthDomain
from geoworkbench.printing.interpretation_report import (
    build_interpretation_report,
    export_interpretation_report_pdf,
    interpretation_report_html,
)
from geoworkbench.printing.interpretation_report_office import (
    export_interpretation_report_docx,
    export_interpretation_report_xlsx,
)
from geoworkbench.printing.report_visual_system import modern_oilfield_report_profile
from geoworkbench.project.cuttings_controller import CuttingsController
from geoworkbench.project.session import ProjectSession
from geoworkbench.project.stratigraphy_controller import StratigraphyController
from geoworkbench.services.interval_gas_statistics import (
    IntervalGasStatisticsIndex,
    build_interval_component_sum_statistics,
    build_interval_statistics,
)
from geoworkbench.services.localization import AppLanguage
from geoworkbench.ui.interpretation_report_dialog import InterpretationReportDialog


def _session() -> ProjectSession:
    session = ProjectSession()
    session.project.name = "Field <Alpha>"
    depth = np.arange(500.0, 521.0, 1.0)
    dataset = Dataset(
        "dataset",
        "Depth log",
        DatasetKind.GTI,
        DepthDomain.MD,
        depth,
    )
    dataset.upsert_curve("TG", np.arange(100.0, 121.0, 1.0), unit="ppm")
    dataset.upsert_curve("C1", np.arange(1.0, 22.0, 1.0), unit="ppm")
    for mnemonic, value in (
        ("C2", 2.0),
        ("C3", 3.0),
        ("IC4", 4.0),
        ("NC4", 5.0),
        ("IC5", 6.0),
        ("NC5", 7.0),
    ):
        dataset.upsert_curve(mnemonic, np.full(depth.shape, value), unit="ppm")
    session.add_dataset(
        dataset,
        "Well 12",
    )
    controller = CuttingsController(session)
    controller.create_full_sample(
        500.0,
        510.0,
        {"sandstone": 70.0, "clay": 30.0},
        calcite_percent=62.5,
        dolomite_percent=17.5,
        lba_group=3,
        lba_type_id="МСБ",
        lba_intensity=4,
        lba_color="yellow-white",
        lba_distribution="Ring",
        lba_cut="Streaming",
        lba_cut_speed="Slow",
        lba_cut_color="Pale yellow",
        lba_residue_type="Oily residue",
        lba_residue_color="Brown",
        lba_odour="Petroleum",
        lba_stain="Present",
        lba_description="bright <direct> fluorescence",
        analysis_interpretation="Manual conclusion\nRequires correlation",
        description="<p>Fine-grained sandstone &amp; clay.</p>",
    )
    controller.create_full_sample(
        510.0,
        520.0,
        {"limestone": 80.0, "marl": 20.0},
        calcite_percent=40.0,
        description="Limestone with marl interbeds.",
    )
    stratigraphy = StratigraphyController(session)
    stratigraphy.add(
        500.0,
        520.0,
        "K1",
        rank="System / Period",
        name="Lower Cretaceous",
        description="Regional stratigraphic interval",
    )
    stratigraphy.add(
        505.0,
        515.0,
        "K1a",
        rank="Formation",
        name="Albian formation",
        description="Target formation",
    )
    return session


def test_interpretation_report_uses_source_results_and_manual_conclusion() -> None:
    report = build_interpretation_report(_session())

    assert report.project_name == "Field <Alpha>"
    assert report.well_name == "Well 12"
    assert report.dataset_name == "Depth log"
    assert report.calcimetry_count == 2
    assert report.lba_count == 1
    assert report.interpreted_count == 1
    assert report.sample_count == 2
    assert len(report.meter_geology) == 20
    assert len(report.stratigraphy) == 2
    first = report.entries[0]
    assert first.insoluble_residue_percent == 20.0
    assert ("intensity", "4") in first.lba_observations
    assert first.lba_standard_assessment is not None
    assert first.lba_standard_assessment.standard.code == "МСБ"
    assert first.interpretation == "Manual conclusion\nRequires correlation"
    gas = {(item.kind, item.mnemonic): item for item in first.gas_statistics}
    assert gas[("total", "TG")].minimum == 100.0
    assert gas[("total", "TG")].mean == 105.0
    assert gas[("total", "TG")].maximum == 110.0
    assert gas[("component", "C1")].minimum == 1.0
    component_sum = gas[("sum", "SUM_COMPONENTS")]
    assert component_sum.minimum == 28.0
    assert component_sum.mean == 33.0
    assert component_sum.maximum == 38.0
    assert [(item.lithotype_id, item.percentage) for item in first.rock_components] == [
        ("sandstone", 70.0),
        ("clay", 30.0),
    ]
    assert first.rock_description == "Fine-grained sandstone & clay."
    assert {item.code for item in first.stratigraphy} == {"K1", "K1a"}
    first_meter = report.meter_geology[0]
    assert (first_meter.top_depth, first_meter.bottom_depth) == (500.0, 501.0)
    assert first_meter.sampling_coverage == 1.0
    assert first_meter.sample_intervals == ((500.0, 510.0),)
    assert [(item.lithotype_id, item.percentage) for item in first_meter.rock_components] == [
        ("sandstone", 70.0),
        ("clay", 30.0),
    ]


def test_meter_geology_preserves_partial_sampling_and_length_weighted_composition() -> None:
    session = ProjectSession()
    session.add_dataset(
        Dataset(
            "partial",
            "Partial sampling",
            DatasetKind.GTI,
            DepthDomain.MD,
            np.array([100.0, 102.0]),
        ),
        "Well partial",
    )
    controller = CuttingsController(session)
    controller.create_full_sample(
        100.0,
        100.25,
        {"sandstone": 100.0},
        description="Upper quarter",
    )
    controller.create_full_sample(
        100.5,
        101.0,
        {"clay": 100.0},
        description="Lower half",
    )

    report = build_interpretation_report(session)

    assert len(report.meter_geology) == 1
    meter = report.meter_geology[0]
    assert (meter.top_depth, meter.bottom_depth) == (100.0, 101.0)
    assert meter.sampling_coverage == 0.75
    assert meter.sample_intervals == ((100.0, 100.25), (100.5, 101.0))
    components = {item.lithotype_id: item.percentage for item in meter.rock_components}
    assert components["sandstone"] == pytest.approx(100.0 / 3.0)
    assert components["clay"] == pytest.approx(200.0 / 3.0)
    assert meter.rock_descriptions == ("Upper quarter", "Lower half")


def test_primary_report_keeps_adjacent_cuttings_samples_separate() -> None:
    session = ProjectSession()
    session.add_dataset(
        Dataset(
            "adjacent",
            "Adjacent samples",
            DatasetKind.GTI,
            DepthDomain.MD,
            np.array([100.0, 103.0]),
        ),
        "Well adjacent",
    )
    controller = CuttingsController(session)
    controller.create_full_sample(
        100.5,
        101.5,
        {"sandstone": 100.0},
        description="Sandstone sample",
    )
    controller.create_full_sample(
        101.5,
        102.5,
        {"clay": 100.0},
        description="Clay sample",
    )

    report = build_interpretation_report(session)
    html = interpretation_report_html(report, AppLanguage.RU)

    actual_start = html.index("Фактические интервалы отбора шлама")
    appendix_start = html.index("Аналитическое приложение: метровая агрегация")
    actual_section = html[actual_start:appendix_start]
    appendix = html[appendix_start:]

    assert "100.5-101.5 m" in actual_section
    assert "101.5-102.5 m" in actual_section
    assert "Песчаник (SANDSTONE): 100%" in actual_section
    assert "Глина (CLAY): 100%" in actual_section
    assert "Песчаник (SANDSTONE): 50%" not in actual_section
    assert "Глина (CLAY): 50%" not in actual_section
    assert "101-102 m" in appendix
    assert "Песчаник (SANDSTONE): 50%" in appendix
    assert "Глина (CLAY): 50%" in appendix


def test_component_sum_is_reported_when_total_gas_curve_is_missing() -> None:
    session = ProjectSession()
    dataset = Dataset(
        "gas-without-total",
        "Components only",
        DatasetKind.GTI,
        DepthDomain.MD,
        np.array([100.0, 101.0, 102.0]),
    )
    dataset.upsert_curve("C1", np.array([1.0, 2.0, 3.0]), unit="ppm")
    dataset.upsert_curve("C2", np.array([10.0, 20.0, 30.0]), unit="ppm")
    session.add_dataset(dataset, "Well components")
    CuttingsController(session).create_full_sample(
        100.0,
        102.0,
        {"sandstone": 100.0},
    )

    report = build_interpretation_report(session)
    entry = report.entries[0]

    assert not any(item.kind == "total" for item in entry.gas_statistics)
    component_sum = next(item for item in entry.gas_statistics if item.kind == "sum")
    assert component_sum.minimum == 11.0
    assert component_sum.mean == 22.0
    assert component_sum.maximum == 33.0
    assert component_sum.unit == "ppm"
    html = interpretation_report_html(report, AppLanguage.RU)
    assert "Сумма компонентов [ppm]" in html
    assert "Total Gas (отдельная кривая)" not in html


def test_component_sum_converts_compatible_mixed_units_and_rejects_incompatible_units() -> None:
    dataset = Dataset(
        "mixed-gas-units",
        "Mixed gas units",
        DatasetKind.GTI,
        DepthDomain.MD,
        np.array([100.0, 101.0]),
    )
    dataset.upsert_curve("C1", np.array([10_000.0, 20_000.0]), unit="ppm")
    c2 = dataset.upsert_curve("C2", np.array([2.0, 3.0]), unit="%vol")

    statistics = build_interval_component_sum_statistics(dataset, 100.0, 101.0)

    assert statistics is not None
    assert statistics.unit == "%vol"
    assert statistics.minimum == 3.0
    assert statistics.mean == 4.0
    assert statistics.maximum == 5.0

    c2.metadata = replace(c2.metadata, unit="kg/m3")
    assert build_interval_component_sum_statistics(dataset, 100.0, 101.0) is None


def test_report_preserves_authored_stratigraphy_names_per_language() -> None:
    session = ProjectSession()
    session.add_dataset(
        Dataset(
            "localized-stratigraphy",
            "Localized stratigraphy",
            DatasetKind.GTI,
            DepthDomain.MD,
            np.array([100.0, 110.0]),
        ),
        "Well localized",
    )
    StratigraphyController(session).add(
        100.0,
        110.0,
        "K",
        rank="System / Period",
        name_i18n={
            "ru": "Меловая система",
            "kk": "Бор жүйесі",
            "en": "Cretaceous System",
        },
        description_i18n={
            "ru": "Русское описание",
            "kk": "Қазақша сипаттама",
            "en": "English description",
        },
    )

    report = build_interpretation_report(session, language=AppLanguage.EN)
    item = report.stratigraphy[0]

    assert item.localized_name(AppLanguage.RU) == "Меловая система"
    assert item.localized_name(AppLanguage.KK) == "Бор жүйесі"
    assert item.localized_name(AppLanguage.EN) == "Cretaceous System"
    assert item.description == "English description"


def test_excel_factual_sample_rows_keep_interval_description_pairing(tmp_path) -> None:
    session = ProjectSession()
    session.add_dataset(
        Dataset(
            "xlsx-description-pairing",
            "XLSX description pairing",
            DatasetKind.GTI,
            DepthDomain.MD,
            np.array([100.0, 103.0]),
        ),
        "Well pairing",
    )
    controller = CuttingsController(session)
    controller.create_full_sample(
        100.5,
        101.5,
        {"sandstone": 100.0},
        description="Описание только первого интервала",
    )
    controller.create_full_sample(
        101.5,
        102.5,
        {"clay": 100.0},
        description="Описание только второго интервала",
    )

    report = build_interpretation_report(session)
    target = export_interpretation_report_xlsx(
        report,
        tmp_path / "interval-description-pairing.xlsx",
        language=AppLanguage.RU,
    )
    workbook = load_workbook(target, read_only=True, data_only=True)
    try:
        rows = list(workbook[workbook.sheetnames[1]].values)
    finally:
        workbook.close()

    header = rows[0]
    interval_column = header.index("Интервал")
    description_column = header.index("Описание пород")
    factual = {
        row[interval_column]: row[description_column]
        for row in rows[1:]
    }

    assert factual == {
        "100.5–101.5 m": "Описание только первого интервала",
        "101.5–102.5 m": "Описание только второго интервала",
    }

def test_interpretation_report_html_is_localized_and_escapes_project_data() -> None:
    report = build_interpretation_report(_session())

    html = interpretation_report_html(report, AppLanguage.RU)
    english = interpretation_report_html(report, AppLanguage.EN)

    assert "Геологический отчёт по шламу, стратиграфии, газу, кальциметрии и ЛБА" in html
    assert "Аналитическое приложение: метровая агрегация" in html
    assert "Фактические интервалы отбора шлама" in html
    assert html.index("Фактические интервалы отбора шлама") < html.index(
        "Аналитическое приложение: метровая агрегация"
    )
    assert "не фактическая шламограмма" in html
    assert "Стратиграфия по всей глубине скважины" in html
    assert "Газ и ЛБА по фактическим интервалам отбора" in html
    assert "Общий газ [ppm]" in html
    assert "Содержание метана [ppm]" in html
    assert "Total Gas (отдельная кривая): TG [ppm]" not in html
    assert "Сумма компонентов [ppm]" in html
    assert "мин 28; среднее 33; макс 38" in html
    assert "не подменяет Total Gas" in html
    assert "отсчётов:" not in html
    assert "Песчаник (SANDSTONE): 70%" in html
    assert "500-510 m" in html
    assert "Lower Cretaceous" in html
    assert "Fine-grained sandstone &amp; clay." in html
    assert "Нерастворимый остаток: 20%" in html
    assert "Интенсивность:</b> 4" in html
    assert "Field &lt;Alpha&gt;" in html
    assert "bright &lt;direct&gt; fluorescence" in html
    assert "Скорость cut:</b> Slow" in html
    assert "Цвет остатка:</b> Brown" in html
    assert "Запах:</b> Petroleum" in html
    assert "Масляное окрашивание:</b> Present" in html
    assert "Оценка по стандарту ЛБА" in html
    assert "МСБ — маслянисто-смолистый битумоид" in html
    assert "html, body { background: #ffffff; color: #172033; }" in html
    assert "td { background: #ffffff; color: #172033; }" in html
    assert "Geological report: cuttings, stratigraphy, gas, calcimetry and LBA" in english
    assert "Analytical appendix: one-metre aggregation" in english
    assert "This report is not an automatic" in english
    assert "Total Gas [ppm]" in english
    assert "Methane [ppm]" in english
    assert "TG [ppm]" not in english
    kazakh = interpretation_report_html(report, AppLanguage.KK)
    assert "автоматты қорытынды болып табылмайды" in kazakh
    assert "Жалпы газ [ppm]" in kazakh
    assert "Метан [ppm]" in kazakh


def test_total_gas_source_alias_is_not_exposed_in_visible_report() -> None:
    report = build_interpretation_report(_session())
    first = report.entries[0]
    total_index = next(
        index
        for index, item in enumerate(first.gas_statistics)
        if item.kind == "total"
    )
    statistics = list(first.gas_statistics)
    statistics[total_index] = replace(
        statistics[total_index],
        mnemonic="TOTAL_GAS_CALC",
    )
    report = replace(
        report,
        entries=(
            replace(first, gas_statistics=tuple(statistics)),
            *report.entries[1:],
        ),
    )

    russian = interpretation_report_html(report, AppLanguage.RU)
    kazakh = interpretation_report_html(report, AppLanguage.KK)
    english = interpretation_report_html(report, AppLanguage.EN)

    assert "Общий газ [ppm]" in russian
    assert "Жалпы газ [ppm]" in kazakh
    assert "Total Gas [ppm]" in english
    assert "TOTAL_GAS_CALC" not in russian
    assert "TOTAL_GAS_CALC" not in kazakh
    assert "TOTAL_GAS_CALC" not in english


def test_interpretation_report_exports_pdf(qapp, tmp_path) -> None:
    report = build_interpretation_report(_session())
    target = tmp_path / "interpretation.pdf"

    exported = export_interpretation_report_pdf(report, target, language=AppLanguage.EN)

    assert exported == target
    assert target.read_bytes().startswith(b"%PDF")
    assert target.stat().st_size > 1000
    with fitz.open(target) as document:
        text = "\n".join(page.get_text() for page in document)
        assert document.page_count >= 3
    assert "Analytical appendix: one-metre aggregation" in text
    assert "Actual cuttings sampling intervals" in text
    assert text.index("Actual cuttings sampling intervals") < text.index(
        "Analytical appendix: one-metre aggregation"
    )
    assert "Whole-well stratigraphy" in text
    assert "Gas and LBA by actual sampling interval" in text
    assert "Total Gas [ppm]" in text
    assert "Methane [ppm]" in text
    assert "TG [ppm]" not in text
    assert "Component sum [ppm]" in text
    assert "Petroleum" in text
    assert "Lower Cretaceous" in text


def test_interpretation_report_exports_excel_and_word(tmp_path) -> None:
    report = build_interpretation_report(_session())
    report = replace(
        report,
        entries=(
            replace(report.entries[0], rock_description="=2+2"),
            *report.entries[1:],
        ),
    )
    xlsx = tmp_path / "geology.xlsx"
    docx = tmp_path / "geology.docx"

    export_interpretation_report_xlsx(
        report, xlsx, language=AppLanguage.RU
    )
    export_interpretation_report_docx(
        report, docx, language=AppLanguage.RU
    )

    workbook = load_workbook(xlsx, data_only=False)
    assert len(workbook.sheetnames) == 5
    summary_values = tuple(
        cell.value
        for row in workbook[workbook.sheetnames[0]].iter_rows()
        for cell in row
    )
    assert "Фактические отборы" in summary_values
    assert "отсчётов" not in summary_values
    assert workbook.sheetnames[1].startswith("Фактические интервалы")
    assert workbook.sheetnames[-1].startswith("Аналитическое приложение")

    samples_sheet = workbook[workbook.sheetnames[1]]
    sample_values = tuple(
        cell.value for row in samples_sheet.iter_rows() for cell in row
    )
    assert "Нерастворимый остаток, %" in sample_values
    assert "Petroleum" in sample_values
    assert "'=2+2" in sample_values

    visual = modern_oilfield_report_profile()
    for table_sheet in workbook.worksheets[1:]:
        assert table_sheet.print_title_rows == "$1:$1"
        assert table_sheet.print_title_cols == "$A:$A"
        assert str(table_sheet.page_setup.paperSize) == str(table_sheet.PAPERSIZE_A4)
        assert table_sheet.page_setup.orientation == table_sheet.ORIENTATION_LANDSCAPE
        assert table_sheet.page_setup.fitToWidth == 0
        assert table_sheet.page_setup.fitToHeight == 0
        assert table_sheet.page_setup.scale == 100
        assert table_sheet.row_dimensions[1].height >= 24.0
        assert table_sheet["A1"].font.sz == pytest.approx(visual.typography.table_pt)
        assert table_sheet["A1"].alignment.wrap_text is True
        assert table_sheet["A1"].alignment.horizontal == "center"
        assert table_sheet["A1"].fill.fgColor.rgb[-6:] == visual.palette.table_header[1:]
        assert table_sheet.oddFooter.left.text == visual.brand_wordmark.replace("&", "&&")
        assert table_sheet.oddFooter.right.text == "&P / &N"
    assert samples_sheet["A2"].fill.fgColor.rgb[-6:] == visual.palette.table_alt[1:]

    gas_sheet = workbook[workbook.sheetnames[2]]
    gas_values = tuple(
        cell.value for row in gas_sheet.iter_rows() for cell in row
    )
    assert "Сумма компонентов" in gas_values
    assert "Общий газ" in gas_values
    assert "Содержание метана" in gas_values
    assert "Mnemonic" not in gas_values
    assert "TG" not in gas_values
    numeric_cells = [
        cell
        for row in gas_sheet.iter_rows(min_row=2, min_col=4, max_col=6)
        for cell in row
        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
    ]
    assert numeric_cells
    assert all(cell.alignment.horizontal == "right" for cell in numeric_cells)

    meter_sheet = workbook[workbook.sheetnames[-1]]
    assert meter_sheet["A1"].comment is not None
    assert "не фактическая шламограмма" in meter_sheet["A1"].comment.text

    with zipfile.ZipFile(docx) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "Фактические интервалы отбора шлама" in document_xml
    assert "Аналитическое приложение: метровая агрегация" in document_xml
    assert document_xml.index("Фактические интервалы отбора шлама") < document_xml.index(
        "Аналитическое приложение: метровая агрегация"
    )
    assert "Сумма компонентов" in document_xml
    assert "Общий газ" in document_xml
    assert "Содержание метана" in document_xml
    assert "Petroleum" in document_xml
    assert "Нерастворимый остаток" in document_xml


def test_total_carbonate_is_exported_as_aggregate_in_excel_and_word(tmp_path) -> None:
    session = _session()
    well = session.current_well
    assert well is not None
    sample = well.cuttings[0]
    sample.calcite_percent = None
    sample.dolomite_percent = None
    sample.total_carbonate_percent = 43.0
    report = build_interpretation_report(session)
    xlsx = export_interpretation_report_xlsx(report, tmp_path / "aggregate.xlsx")
    docx = export_interpretation_report_docx(report, tmp_path / "aggregate.docx")

    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        rows = list(workbook[workbook.sheetnames[1]].values)
        assert "Общая карбонатность, %" in rows[0]
        assert rows[1][rows[0].index("Общая карбонатность, %")] == 43.0
        assert rows[1][rows[0].index("CaCO3, %")] is None
    finally:
        workbook.close()
    with zipfile.ZipFile(docx) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "Общая карбонатность: 43%" in document_xml
    assert "CaCO3: 43%" not in document_xml




def test_interval_gas_statistics_index_matches_one_shot_helpers() -> None:
    session = _session()
    dataset = session.current_dataset
    assert dataset is not None
    index = IntervalGasStatisticsIndex(dataset)

    for top, bottom in ((500.0, 510.0), (505.0, 515.0), (515.0, 505.0)):
        assert index.build(top, bottom) == build_interval_statistics(
            dataset, top, bottom
        )
        assert index.build_component_sum(
            top, bottom
        ) == build_interval_component_sum_statistics(dataset, top, bottom)


def test_interpretation_report_dialog_previews_report(qapp) -> None:
    dialog = InterpretationReportDialog(_session(), language=AppLanguage.EN)
    dialog.resize(720, 420)
    dark_palette = QPalette(dialog.palette())
    dark_palette.setColor(QPalette.ColorRole.Window, QColor("#252a31"))
    dark_palette.setColor(QPalette.ColorRole.Base, QColor("#30363d"))
    dark_palette.setColor(QPalette.ColorRole.Text, QColor("#e5e7eb"))
    dialog.setPalette(dark_palette)
    dialog.show()
    qapp.processEvents()

    preview = dialog.findChild(QTextBrowser, "interpretation-report-preview")
    export_button = dialog.findChild(QPushButton, "interpretation-report-export")
    xlsx_button = dialog.findChild(
        QPushButton, "interpretation-report-export-xlsx"
    )
    docx_button = dialog.findChild(
        QPushButton, "interpretation-report-export-docx"
    )

    assert preview is not None
    assert "Manual conclusion" in preview.toPlainText()
    assert "background-color: #ffffff" in preview.styleSheet()
    assert "color: #172033" in preview.styleSheet()
    assert (
        preview.horizontalScrollBarPolicy()
        is Qt.ScrollBarPolicy.ScrollBarAlwaysOn
    )
    assert preview.verticalScrollBarPolicy() is Qt.ScrollBarPolicy.ScrollBarAlwaysOn
    assert preview.lineWrapMode() is QTextEdit.LineWrapMode.FixedPixelWidth
    assert preview.lineWrapColumnOrWidth() == 1600
    assert preview.horizontalScrollBar().maximum() > 0
    assert preview.verticalScrollBar().maximum() > 0
    viewport = preview.viewport()
    image = viewport.grab().toImage()
    canvas_pixel = image.pixelColor(max(0, image.width() - 8), max(0, image.height() - 8))
    assert canvas_pixel.lightness() > 220
    assert export_button is not None and export_button.text() == "Export PDF..."
    assert xlsx_button is not None and xlsx_button.text() == "Export Excel..."
    assert docx_button is not None and docx_button.text() == "Export Word..."
    dialog.close()
