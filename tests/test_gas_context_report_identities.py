from copy import deepcopy
from html import escape
import xml.etree.ElementTree as ET
import zipfile
from unicodedata import normalize

import fitz
import numpy as np
from openpyxl import load_workbook
import pytest

from geoworkbench.data.hydrocarbon_interpretation_export import (
    export_hydrocarbon_interpretation_docx, export_hydrocarbon_interpretation_xlsx,
)
from geoworkbench.data.hydrocarbon_interpretation_export_docx_polished import export_polished_hydrocarbon_interpretation_docx
from geoworkbench.domain.gas_context_events import GasContextEvent, GasContextEventType, InterpretationImpact
from geoworkbench.domain.models import CurveData, CurveMetadata, DepthDomain
from geoworkbench.printing.hydrocarbon_interpretation_report import export_hydrocarbon_interpretation_pdf
from geoworkbench.project.controller import ProjectController
from geoworkbench.services.gas_context_report_labels import (
    gas_context_event_label, gas_context_type_text,
    gas_context_impact_label, gas_context_identity_label,
)
from geoworkbench.services.hydrocarbon_interpretation import (
    build_hydrocarbon_interpretation_report, build_opus_interpretation_report, hydrocarbon_interpretation_html,
)
from geoworkbench.services.localization import AppLanguage
from test_hydrocarbon_interpretation import _session


@pytest.mark.parametrize("language", list(AppLanguage))
@pytest.mark.parametrize("impact", list(InterpretationImpact))
def test_all_impacts_have_human_readable_localized_labels(language, impact):
    label = gas_context_impact_label(impact, language)
    assert label and label != impact.value
    assert "_" not in label


@pytest.mark.parametrize("language", list(AppLanguage))
@pytest.mark.parametrize("builder", [build_hydrocarbon_interpretation_report, build_opus_interpretation_report])
def test_reopened_context_identity_is_consistent_in_all_client_exports(qapp, tmp_path, language, builder):
    session = _session()
    dataset = session.current_dataset
    dataset.curves["TG"] = CurveData(CurveMetadata("TG", "TG", "TG", "%", "Measured TG", dataset.dataset_id), np.full(dataset.depth.shape, 4.0))
    events = [
        GasContextEvent("connection-one", GasContextEventType.CONNECTION_GAS, 1039, 1043,
                        depth_domain=DepthDomain.MD, reported_total_gas=4.25, reported_unit="%"),
        GasContextEvent("connection-two", GasContextEventType.CONNECTION_GAS, 1050, 1055, depth_domain=DepthDomain.MD),
        GasContextEvent("formation-one", GasContextEventType.FORMATION_SHOW, 1060, 1065, depth_domain=DepthDomain.MD),
        GasContextEvent("review-one", GasContextEventType.ELEVATED_UNCLASSIFIED, 1070, 1075, depth_domain=DepthDomain.MD),
        GasContextEvent("=SUM(1,1)<&>", GasContextEventType.SWAB_GAS, 1080, 1085, depth_domain=DepthDomain.MD),
        GasContextEvent("draft-only", GasContextEventType.TRIP_GAS, 1020, 1025, depth_domain=DepthDomain.MD, confirmed=False),
        GasContextEvent("excluded-only", GasContextEventType.CALIBRATION_GAS, 1090, 1095, depth_domain=DepthDomain.MD),
    ]
    session.current_well.gas_context_events.extend(events)
    package = tmp_path / "context.geologpkg"
    ProjectController(session=session).save_project(package)
    reopened = ProjectController().open_project(package)
    dataset = reopened.current_dataset
    before_arrays = {key: curve.values.copy() for key, curve in dataset.curves.items()}
    before_registry = deepcopy(reopened.current_well.gas_context_events)
    report = builder(reopened)
    before_report = deepcopy(report)
    visible = report.gas_context_events
    assert len(visible) == 5
    html = hydrocarbon_interpretation_html(report, language)
    for event in visible:
        assert escape(event.event_id) in html
        assert gas_context_type_text(event.event_type, language) in html
        assert gas_context_impact_label(event.effective_impact, language) in html
    assert "<td>=SUM(1,1)<&>" not in html
    assert "draft-only" not in html and "excluded-only" not in html
    assert gas_context_identity_label(language) in html

    for export in (export_hydrocarbon_interpretation_docx, export_polished_hydrocarbon_interpretation_docx):
        target = tmp_path / (export.__name__ + ".docx")
        export(report, target, dataset=dataset, language=language)
        with zipfile.ZipFile(target) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
            text = "\n".join(e.text or "" for e in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
        for event in visible:
            assert event.event_id in text
            assert gas_context_type_text(event.event_type, language) in text
            assert gas_context_impact_label(event.effective_impact, language) in text
        assert "draft-only" not in text and "excluded-only" not in text
        assert gas_context_identity_label(language) in text

    xlsx = export_hydrocarbon_interpretation_xlsx(report, dataset, tmp_path / "context.xlsx", language=language)
    workbook = load_workbook(xlsx, data_only=False)
    try:
        title = {AppLanguage.RU: "Газовый контекст", AppLanguage.KK: "Газ контексті", AppLanguage.EN: "Gas context"}[language]
        sheet = workbook[title]
        assert sheet["Q1"].value == gas_context_identity_label(language)
        for row, event in enumerate(visible, 2):
            assert sheet.cell(row, 1).value == gas_context_type_text(event.event_type, language)
            assert sheet.cell(row, 5).value == gas_context_impact_label(event.effective_impact, language)
            assert sheet.cell(row, 17).value.lstrip("'") == event.event_id
            assert sheet.cell(row, 17).data_type == "s"
            assert sheet.cell(row, 2).value == event.top_depth
            assert sheet.cell(row, 3).value == event.bottom_depth
        assert sheet["G2"].value == 4.0 and sheet["K2"].value == 4.25
        assert sheet["G2"].data_type == "n" and sheet["K2"].data_type == "n"
    finally:
        workbook.close()

    target = tmp_path / "context.pdf"
    export_hydrocarbon_interpretation_pdf(report, target, dataset=dataset, language=language)
    with fitz.open(target) as document:
        text = "\n".join(page.get_text() for page in document)
        compact = "".join(normalize("NFKC", text).split())
        for event in visible:
            assert "".join(event.event_id.split()) in compact
            assert "".join(gas_context_event_label(event.event_type, language).split()) in compact
            assert "".join(gas_context_impact_label(event.effective_impact, language).split()) in compact
        assert "draft-only" not in text and "excluded-only" not in text
        assert gas_context_identity_label(language) in text
    assert reopened.current_well.gas_context_events == before_registry
    assert report.gas_context_events == before_report.gas_context_events
    assert report.candidates == before_report.candidates
    assert report.suppressed_candidates == before_report.suppressed_candidates
    for key, values in before_arrays.items():
        np.testing.assert_array_equal(dataset.curves[key].values, values)
