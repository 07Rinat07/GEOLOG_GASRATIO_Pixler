"""Opt-in Qt comparison handoff to the existing readable Excel export."""

from __future__ import annotations

from dataclasses import replace
import faulthandler
import os
from pathlib import Path

import pytest

from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QMessageBox
from openpyxl import load_workbook  # type: ignore[import-untyped]

from geoworkbench.data import hydrocarbon_interpretation_export_readable as readable
from geoworkbench.services.interpretation_method_comparison_html import (
    method_comparison_table,
)
from geoworkbench.services.localization import AppLanguage
from geoworkbench.ui import interpretation_report_workspace_legacy as legacy
from test_interpretation_method_comparison_chart_html import _comparison
from test_interpretation_method_comparison_explicit_ui import _workspace


@pytest.fixture(autouse=True)
def _bound_gui_acceptance():
    faulthandler.dump_traceback_later(35, exit=True)
    try:
        yield
    finally:
        faulthandler.cancel_dump_traceback_later()


@pytest.fixture(autouse=True)
def _no_modal_gui_dialogs(monkeypatch: pytest.MonkeyPatch) -> None:
    def reject(_parent: object, title: str, message: str, *args: object, **kwargs: object) -> None:
        del args, kwargs
        pytest.fail(f"Unexpected blocking XLSX Qt dialog {title!r}: {message}")

    for method in ("critical", "warning", "information", "question"):
        monkeypatch.setattr(QMessageBox, method, reject)


def _ci_stage(stage: str) -> None:
    if not os.environ.get("CI"):
        return
    path = Path("build/ci-artifacts/quality/ui-xlsx-stage.txt")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as output:
        output.write(stage + "\n")


def _assert_explicit_comparison_included_in_real_readable_xlsx(
    tmp_path: Path,
    qapp,
    monkeypatch: pytest.MonkeyPatch,
    language: AppLanguage,
    sheet_name: str,
) -> None:
    _ci_stage(f"xlsx start {language.value}")
    workspace = _workspace(language)
    _ci_stage("workspace constructed")
    assert workspace.report is not None

    def opus(session, **kwargs):
        del session, kwargs
        assert workspace.report is not None
        return replace(workspace.report, report_profile="opus")

    monkeypatch.setattr(legacy, "build_opus_interpretation_report", opus)
    monkeypatch.setattr(
        legacy, "compare_standard_opus_reports",
        lambda standard, companion: (_comparison(),),
    )
    target = tmp_path / "explicit.xlsx"
    monkeypatch.setattr(workspace, "_choose_target", lambda suffix, file_filter: target)
    monkeypatch.setattr(workspace, "_show_export_success", lambda path: None)
    _ci_stage("before compare click")
    workspace.compare_methods_button.click()
    qapp.processEvents()
    _ci_stage("after compare click")

    _ci_stage("before xlsx export")
    workspace._export_xlsx()
    _ci_stage("after xlsx export")
    assert target.exists()
    expected = method_comparison_table(
        (_comparison(),), language, depth_unit=workspace.report.depth_unit,
    )
    _ci_stage("before workbook load")
    workbook = load_workbook(target, read_only=False, data_only=False)
    _ci_stage("after workbook load")
    try:
        assert sheet_name in workbook.sheetnames
        sheet = workbook[sheet_name]
        assert tuple(sheet.cell(1, i).value for i in range(1, 6)) == expected.headings
        assert tuple(sheet.cell(2, i).value for i in range(1, 6)) == expected.rows[0].cells
        assert sheet["F2"].value == expected.rows[0].reason_code
        assert sheet["G2"].value == expected.rows[0].reason_version
        assert sheet.column_dimensions["F"].hidden
        assert sheet.column_dimensions["G"].hidden
    finally:
        workbook.close()
    assert workspace._method_comparisons == (_comparison(),)
    _ci_stage("before workspace close")
    workspace.close()
    _ci_stage("after workspace close")


def _assert_default_ui_workbook_keeps_legacy_export_arguments(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, language: AppLanguage,
) -> None:
    workspace = _workspace(language)
    target = tmp_path / "plain.xlsx"
    calls: list[dict[str, object]] = []

    def capture(report, dataset, target, **kwargs):
        del report, dataset
        calls.append(kwargs)
        return target

    monkeypatch.setattr(readable, "export_readable_hydrocarbon_interpretation_xlsx", capture)
    monkeypatch.setattr(workspace, "_choose_target", lambda suffix, file_filter: target)
    monkeypatch.setattr(workspace, "_show_export_success", lambda path: None)
    workspace._export_xlsx()
    assert len(calls) == 1
    assert "method_comparisons" not in calls[0]
    assert calls[0]["language"] is language
    assert calls[0]["identity"] is not None

    assert calls[0]["progress"] == workspace._update_report_export_progress
    workspace.close()


def _assert_stale_pair_never_reaches_xlsx_export(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, change: str,
) -> None:
    workspace = _workspace()
    assert workspace.report is not None

    def opus(session, **kwargs):
        del session, kwargs
        assert workspace.report is not None
        return replace(workspace.report, report_profile="opus")

    monkeypatch.setattr(legacy, "build_opus_interpretation_report", opus)
    monkeypatch.setattr(
        legacy, "compare_standard_opus_reports",
        lambda standard, companion: (_comparison(),),
    )
    workspace.compare_methods_button.click()
    qapp.processEvents()
    assert workspace._method_comparisons

    if change == "report":
        workspace.report = replace(workspace.report, generated_at="new-revision")
    elif change == "dataset":
        dataset = workspace.controller.session.current_dataset
        assert dataset is not None
        dataset.dataset_id = "new-source-id"
    elif change == "curve_revision":
        dataset = workspace.controller.session.current_dataset
        assert dataset is not None
        next(iter(dataset.curves.values())).version += 1
    elif change == "axis_values":
        dataset = workspace.controller.session.current_dataset
        assert dataset is not None
        dataset.active_index.values = dataset.active_index.values.copy()
    else:
        workspace.threshold.blockSignals(True)
        workspace.threshold.setValue(workspace.threshold.value() + 0.5)
        workspace.threshold.blockSignals(False)

    target = tmp_path / "stale.xlsx"
    calls: list[dict[str, object]] = []

    def capture(report, dataset, target, **kwargs):
        del report, dataset
        calls.append(kwargs)
        return target

    monkeypatch.setattr(readable, "export_readable_hydrocarbon_interpretation_xlsx", capture)
    monkeypatch.setattr(workspace, "_choose_target", lambda suffix, file_filter: target)
    monkeypatch.setattr(workspace, "_show_export_success", lambda path: None)
    workspace._export_xlsx()
    assert len(calls) == 1
    assert "method_comparisons" not in calls[0]
    assert workspace._method_comparisons == ()
    assert workspace._comparison_source_key is None
    workspace.close()


# One real Qt workspace per test process on the Windows native-heavy runner.
# Preserve all nine RU/KK/EN, legacy-kwargs and stale-input assertions.
# Native 0xC0000005 remains tracked independently as issue #544.


def test_explicit_comparison_included_in_real_readable_xlsx_ru(
    tmp_path: Path, qapp, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_explicit_comparison_included_in_real_readable_xlsx(
        tmp_path, qapp, monkeypatch, AppLanguage.RU, "Сравнение методов",
    )


def test_explicit_comparison_included_in_real_readable_xlsx_kk(
    tmp_path: Path, qapp, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_explicit_comparison_included_in_real_readable_xlsx(
        tmp_path, qapp, monkeypatch, AppLanguage.KK, "Әдістерді салыстыру",
    )


def test_explicit_comparison_included_in_real_readable_xlsx_en(
    tmp_path: Path, qapp, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_explicit_comparison_included_in_real_readable_xlsx(
        tmp_path, qapp, monkeypatch, AppLanguage.EN, "Method comparison",
    )


def test_default_ui_workbook_keeps_legacy_export_arguments_ru(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_default_ui_workbook_keeps_legacy_export_arguments(
        qapp, tmp_path, monkeypatch, AppLanguage.RU,
    )


def test_default_ui_workbook_keeps_legacy_export_arguments_kk(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_default_ui_workbook_keeps_legacy_export_arguments(
        qapp, tmp_path, monkeypatch, AppLanguage.KK,
    )


def test_default_ui_workbook_keeps_legacy_export_arguments_en(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_default_ui_workbook_keeps_legacy_export_arguments(
        qapp, tmp_path, monkeypatch, AppLanguage.EN,
    )


def test_stale_pair_never_reaches_xlsx_export_report(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_stale_pair_never_reaches_xlsx_export(
        qapp, tmp_path, monkeypatch, "report",
    )


def test_stale_pair_never_reaches_xlsx_export_dataset(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_stale_pair_never_reaches_xlsx_export(
        qapp, tmp_path, monkeypatch, "dataset",
    )


def test_stale_pair_never_reaches_xlsx_export_threshold(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_stale_pair_never_reaches_xlsx_export(
        qapp, tmp_path, monkeypatch, "threshold",
    )




def _assert_queued_change_drops_comparison(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, change: str,
) -> None:
    """A queued Qt event must be handled before we capture the Excel snapshot."""

    workspace = _workspace()
    assert workspace.report is not None

    def opus(session, **kwargs):
        del session, kwargs
        assert workspace.report is not None
        return replace(workspace.report, report_profile="opus")

    monkeypatch.setattr(legacy, "build_opus_interpretation_report", opus)
    monkeypatch.setattr(
        legacy, "compare_standard_opus_reports",
        lambda standard, companion: (_comparison(),),
    )
    workspace.compare_methods_button.click()
    qapp.processEvents()
    assert workspace._method_comparisons == (_comparison(),)

    def mutate_source() -> None:
        if change == "threshold":
            workspace.threshold.setValue(workspace.threshold.value() + 0.5)
        elif change == "language":
            index = workspace.report_output_language.findData(AppLanguage.EN)
            assert index >= 0
            workspace.report_output_language.setCurrentIndex(index)
        else:
            assert workspace.report is not None
            workspace.report = replace(workspace.report, generated_at="queued-revision")

    target = tmp_path / "queued.xlsx"
    calls: list[dict[str, object]] = []

    def capture(report, dataset, target, **kwargs):
        del report, dataset
        calls.append(kwargs)
        return target

    monkeypatch.setattr(readable, "export_readable_hydrocarbon_interpretation_xlsx", capture)
    monkeypatch.setattr(workspace, "_choose_target", lambda suffix, file_filter: target)
    monkeypatch.setattr(workspace, "_show_export_success", lambda path: None)
    errors: list[Exception] = []
    monkeypatch.setattr(workspace, "_show_export_error", errors.append)
    # _export_xlsx must not freeze comparison before the progress context
    # delivers the event; no test runner skips or UI signal suppression.
    QTimer.singleShot(0, mutate_source)
    workspace._export_xlsx()
    if change == "report":
        # A report replaced mid-export must not produce even a legacy workbook
        # from the obsolete report object.
        assert calls == []
        assert len(errors) == 1
        assert "source changed" in str(errors[0])
    else:
        assert errors == []
        assert len(calls) == 1
        assert "method_comparisons" not in calls[0]
        if change == "language":
            assert calls[0]["language"] is AppLanguage.EN
            # The passport uses the newly selected output language as well.
            assert calls[0]["identity"] is not None
    assert workspace._method_comparisons == ()
    assert workspace._comparison_source_key is None
    workspace.close()


def test_queued_threshold_change_drops_comparison_before_xlsx_writer(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_queued_change_drops_comparison(qapp, tmp_path, monkeypatch, "threshold")


def test_queued_report_replacement_drops_comparison_before_xlsx_writer(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_queued_change_drops_comparison(qapp, tmp_path, monkeypatch, "report")



def test_queued_language_change_uses_current_xlsx_language_without_comparison(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _assert_queued_change_drops_comparison(qapp, tmp_path, monkeypatch, "language")



def test_xlsx_progress_rejects_queued_comparison_change_before_atomic_save(
    qapp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Do not commit a once-verified comparison after progress pumps Qt events."""

    workspace = _workspace()
    assert workspace.report is not None

    def opus(session, **kwargs):
        del session, kwargs
        assert workspace.report is not None
        return replace(workspace.report, report_profile="opus")

    monkeypatch.setattr(legacy, "build_opus_interpretation_report", opus)
    monkeypatch.setattr(
        legacy, "compare_standard_opus_reports",
        lambda standard, companion: (_comparison(),),
    )
    workspace.compare_methods_button.click()
    qapp.processEvents()
    assert workspace._method_comparisons == (_comparison(),)

    target = tmp_path / "must-not-be-written.xlsx"
    errors: list[Exception] = []
    successes: list[Path] = []

    def simulate_export(report, dataset, path, **kwargs):
        del report, dataset, path
        assert kwargs["method_comparisons"] == (_comparison(),)
        callback = kwargs["progress"]
        assert callable(callback)
        QTimer.singleShot(
            0, lambda: workspace.threshold.setValue(workspace.threshold.value() + 0.5),
        )
        callback("Writing", 20, 100)
        pytest.fail("The stale source must raise before the workbook is saved")

    monkeypatch.setattr(readable, "export_readable_hydrocarbon_interpretation_xlsx", simulate_export)
    monkeypatch.setattr(workspace, "_choose_target", lambda suffix, file_filter: target)
    monkeypatch.setattr(workspace, "_show_export_success", successes.append)
    monkeypatch.setattr(workspace, "_show_export_error", errors.append)
    workspace._export_xlsx()
    assert len(errors) == 1
    assert "comparison changed" in str(errors[0])
    assert successes == []
    assert not target.exists()
    assert workspace._method_comparisons == ()
    workspace.close()


def test_changed_curve_revision_drops_comparison_before_xlsx_writer(qapp, tmp_path, monkeypatch):
    _assert_stale_pair_never_reaches_xlsx_export(qapp, tmp_path, monkeypatch, "curve_revision")


def test_replaced_axis_values_drop_comparison_before_xlsx_writer(qapp, tmp_path, monkeypatch):
    _assert_stale_pair_never_reaches_xlsx_export(qapp, tmp_path, monkeypatch, "axis_values")


def test_real_xlsx_preserves_existing_file_when_curve_changes_before_save(
    qapp, tmp_path, monkeypatch,
):
    workspace = _workspace()
    assert workspace.report is not None
    monkeypatch.setattr(
        legacy, "build_opus_interpretation_report",
        lambda session, **kwargs: replace(workspace.report, report_profile="opus"),
    )
    monkeypatch.setattr(
        legacy, "compare_standard_opus_reports", lambda standard, opus: (_comparison(),),
    )
    workspace.compare_methods_button.click()
    qapp.processEvents()
    assert workspace._method_comparisons == (_comparison(),)
    dataset = workspace.controller.session.current_dataset
    assert dataset is not None
    curve = next(iter(dataset.curves.values()))
    original_values = curve.values.copy()
    target = tmp_path / "existing.xlsx"
    original_bytes = b"existing workbook must remain unchanged"
    target.write_bytes(original_bytes)
    errors = []
    successes = []
    original_progress = workspace._update_report_export_progress
    changed = []

    def progress(stage, current, total):
        if current == 95 and not changed:
            changed.append(True)
            QTimer.singleShot(0, lambda: setattr(curve, "version", curve.version + 1))
        original_progress(stage, current, total)

    monkeypatch.setattr(workspace, "_update_report_export_progress", progress)
    monkeypatch.setattr(workspace, "_choose_target", lambda suffix, file_filter: target)
    monkeypatch.setattr(workspace, "_show_export_error", errors.append)
    monkeypatch.setattr(workspace, "_show_export_success", successes.append)
    try:
        workspace._export_xlsx()
        assert changed == [True]
        assert len(errors) == 1
        assert successes == []
        assert target.read_bytes() == original_bytes
        assert workspace._method_comparisons == ()
        assert list(tmp_path.iterdir()) == [target]
        import numpy as np

        np.testing.assert_array_equal(curve.values, original_values)
    finally:
        workspace.close()
