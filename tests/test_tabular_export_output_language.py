"""Selected interval and statistics use one persisted output language."""

import csv
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QFileDialog

from geoworkbench.data.selection_export import export_selection_text
from geoworkbench.domain.models import (
    CurveData, CurveMetadata, Dataset, DatasetKind, DepthDomain, Project, Well,
)
from geoworkbench.project.session import ProjectSession
from geoworkbench.services.localization import (
    AppLanguage, LanguageSettings, Localizer, TabularExportLanguageSettings,
)
from geoworkbench.services.user_profiles import UserProfileSettings
from geoworkbench.ui.main_window import MainWindow


def _dataset() -> Dataset:
    dataset = Dataset(
        "dataset-c1", "Well-A", DatasetKind.GTI, DepthDomain.MD,
        np.array([100.0, 101.0]),
    )
    dataset.curves["curve-c1"] = CurveData(
        CurveMetadata("curve-c1", "C1", "C1", "%", None, dataset.dataset_id),
        np.array([2.0, 3.0]),
    )
    return dataset


@pytest.mark.parametrize("language", list(AppLanguage))
def test_csv_headers_are_localized_but_measurements_are_not(
    tmp_path: Path, language: AppLanguage,
) -> None:
    dataset = _dataset()
    original_values = dataset.curves["curve-c1"].values.copy()
    path = export_selection_text(
        dataset, tmp_path / "selection.csv", ["curve-c1"], 100.0, 101.0,
        delimiter=",", language=language,
    )
    with path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.reader(stream))
    assert "DEPTH [m]" in rows[0][0]
    assert "C1 [%]" in rows[0][1]
    assert {
        AppLanguage.RU: "Глубина по стволу",
        AppLanguage.KK: "Оқпан бойынша тереңдік",
        AppLanguage.EN: "Measured depth",
    }[language] in rows[0][0]
    assert {
        AppLanguage.RU: "Метан",
        AppLanguage.KK: "Метан",
        AppLanguage.EN: "Methane",
    }[language] in rows[0][1]
    assert rows[1:] == [["100", "2"], ["101", "3"]]
    np.testing.assert_array_equal(dataset.curves["curve-c1"].values, original_values)


def test_language_preference_fallback_and_persistence(tmp_path: Path) -> None:
    settings = QSettings(str(tmp_path / "preferences.ini"), QSettings.Format.IniFormat)
    preference = TabularExportLanguageSettings(settings)
    assert preference.current(AppLanguage.KK) is AppLanguage.KK
    settings.setValue(preference.KEY, "invalid")
    assert preference.current(AppLanguage.RU) is AppLanguage.RU
    preference.save(AppLanguage.EN)
    assert preference.current(AppLanguage.RU) is AppLanguage.EN


@pytest.mark.parametrize(
    ("ui_language", "output_language"),
    [
        (AppLanguage.RU, AppLanguage.KK),
        (AppLanguage.EN, AppLanguage.RU),
        (AppLanguage.KK, AppLanguage.EN),
    ],
)
def test_all_six_outputs_share_language_without_relocalizing_operator_ui(
    qapp, tmp_path: Path, monkeypatch,
    ui_language: AppLanguage, output_language: AppLanguage,
) -> None:
    settings = QSettings(str(tmp_path / "preferences.ini"), QSettings.Format.IniFormat)
    def open_window(language: AppLanguage) -> MainWindow:
        return MainWindow(
            language=language,
            language_settings=LanguageSettings(settings),
            user_profile_settings=UserProfileSettings(settings),
        )

    window = open_window(ui_language)
    dataset = _dataset()
    well = Well("well", "Well", datasets={dataset.dataset_id: dataset})
    session = ProjectSession(
        project=Project("project", "Project", wells={well.well_id: well}),
        current_well_id=well.well_id, current_dataset_id=dataset.dataset_id,
    )
    window.project_controller.session = session
    window._bind_project_session()
    window.dataset_selection.select(dataset, 100.0, 101.0, ("curve-c1",))
    window.tabular_export_language_actions[output_language].trigger()
    assert window.language is ui_language
    assert window.tabular_export_language is output_language
    monkeypatch.setattr(
        window, "_confirm_export_overwrite", lambda *args, **kwargs: False
    )

    outputs: dict[str, Path] = {}
    for extension, handler in (
        ("csv", window.export_selected_csv),
        ("xlsx", window.export_selected_excel),
        ("docx", window.export_selected_docx),
        ("html", window.export_selected_html),
    ):
        target = tmp_path / f"selected.{extension}"
        monkeypatch.setattr(
            QFileDialog, "getSaveFileName",
            lambda *args, path=target, **kwargs: (str(path), ""),
        )
        handler()
        assert target.is_file(), extension
        outputs[extension] = target

    with outputs["csv"].open(encoding="utf-8", newline="") as stream:
        rows = list(csv.reader(stream))
    assert "C1 [%]" in rows[0][1]
    assert rows[1:] == [["100", "2"], ["101", "3"]]
    workbook = load_workbook(outputs["xlsx"])
    assert workbook["Metadata"]["B6"].value == output_language.value
    assert workbook["Data"]["B2"].value == 2.0

    # Selection reports have an authored title, not a generic "Engineering report" title.
    # Assert actual localized section labels in both rendered document formats.
    metadata_heading = {
        AppLanguage.RU: "Параметры отчёта",
        AppLanguage.KK: "Есеп параметрлері",
        AppLanguage.EN: "Report parameters",
    }[output_language]
    html = outputs["html"].read_text(encoding="utf-8")
    assert f'<html lang="{output_language.value}">' in html
    assert "<h1>Well-A selection</h1>" in html
    assert f"<h2>{metadata_heading}</h2>" in html
    with ZipFile(outputs["docx"]) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "Well-A selection" in document_xml
    assert metadata_heading in document_xml

    window._show_interval_analysis_from_gesture({
        "top": 100.0, "bottom": 101.0,
        "axis_id": dataset.active_index_id,
        "axis_label": "Depth", "axis_unit": "m",
        "axis_is_datetime": False, "mnemonics": ("C1",),
    })
    panel = window.interval_statistics_panel
    assert panel.statistics
    original_names = panel.display_names
    original_rows = panel.table.rowCount()
    for extension in ("csv", "xlsx"):
        target = tmp_path / f"stats.{extension}"
        monkeypatch.setattr(
            QFileDialog, "getSaveFileName",
            lambda *args, path=target, **kwargs: (str(path), ""),
        )
        window._export_interval_statistics(extension)
        assert target.is_file()
        parameter = Localizer.create(output_language).text("statistics.parameter")
        if extension == "csv":
            assert parameter in target.read_text(encoding="utf-8-sig")
        else:
            sheet = load_workbook(target).active
            assert sheet is not None
            assert sheet["A4"].value == parameter

    assert panel._language is ui_language
    assert panel.table.rowCount() == original_rows
    assert panel.display_names == original_names
    window.change_language(AppLanguage.RU if ui_language is not AppLanguage.RU else AppLanguage.EN)
    assert window.tabular_export_language is output_language
    window.close()

    reopened = open_window(ui_language)
    assert reopened.tabular_export_language is output_language
    assert reopened.tabular_export_language_actions[output_language].isChecked()
    reopened.close()
