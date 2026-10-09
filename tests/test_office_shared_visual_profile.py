from copy import deepcopy
from dataclasses import replace
import xml.etree.ElementTree as ET
import zipfile

import numpy as np
from openpyxl import load_workbook
import pytest

from geoworkbench.data import hydrocarbon_interpretation_export as shared
from geoworkbench.data import report_document_export as generic
from geoworkbench.printing import interpretation_report_office as geology
from geoworkbench.printing.interpretation_report import build_interpretation_report
from geoworkbench.printing.report_visual_system import modern_oilfield_report_profile
from geoworkbench.project.controller import ProjectController
from geoworkbench.project.session import ProjectSession
from geoworkbench.services.localization import AppLanguage
from geoworkbench.services.report_definition import ReportIntervalContext, resolve_report_definition
from test_interpretation_report import _session
from test_report_document_export import _resolved_report


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{" + NS["w"] + "}"


def _profile(kind, grayscale):
    visual = modern_oilfield_report_profile(grayscale=grayscale)
    if kind == "default":
        return visual
    if kind == "custom":
        return replace(visual, typography=replace(visual.typography, title_pt=24, body_pt=11,
                       table_pt=9.5, caption_pt=6.5, footer_pt=6),
                       layout=replace(visual.layout, thin_rule_pt=0.4, strong_rule_pt=1.2))
    return replace(visual, typography=replace(visual.typography, title_pt=36, body_pt=24,
                   table_pt=20, caption_pt=18, footer_pt=16))


def _docx(target):
    with zipfile.ZipFile(target) as package:
        document = ET.fromstring(package.read("word/document.xml"))
        styles = ET.fromstring(package.read("word/styles.xml"))
    return document, styles


@pytest.mark.parametrize("language", list(AppLanguage))
@pytest.mark.parametrize("profile", ["default", "custom", "large"])
@pytest.mark.parametrize("grayscale", [False, True])
@pytest.mark.parametrize("kind", ["generic_docx", "geology_docx", "geology_xlsx"])
def test_office_profile_survives_production_reopen_without_changing_values(
    tmp_path, monkeypatch, language, profile, grayscale, kind
):
    visual = _profile(profile, grayscale)
    if kind == "generic_docx":
        dataset, resolved = _resolved_report()
        session = ProjectSession()
        session.add_dataset(dataset, "Well Әғқң <&>")
    else:
        session = _session()
    target = tmp_path / "office.geologpkg"
    ProjectController(session=session).save_project(target)
    restored = ProjectController().open_project(target)
    dataset = restored.current_dataset
    source = {key: curve.values.copy() for key, curve in dataset.curves.items()}
    before = deepcopy((restored.current_well.cuttings, restored.current_well.stratigraphy))
    if kind == "generic_docx":
        resolved = resolve_report_definition(dataset, resolved.definition,
                    context=ReportIntervalContext(selection_range=(100, 102)), require_curves=True)
        def export(path):
            return generic.export_report_docx(dataset, path, resolved, language=language)
    else:
        report = build_interpretation_report(restored, language)
        def export(path):
            function = geology.export_interpretation_report_docx if kind == "geology_docx" else geology.export_interpretation_report_xlsx
            return function(report, path, language=language)
    suffix = ".xlsx" if kind == "geology_xlsx" else ".docx"
    baseline = export(tmp_path / ("baseline" + suffix))
    for module in (generic, shared, geology):
        monkeypatch.setattr(module, "modern_oilfield_report_profile", lambda: visual)
    output = export(tmp_path / ("profile" + suffix))
    if kind.endswith("docx"):
        document, styles = _docx(output)
        old, _ = _docx(baseline)
        assert [node.text for node in document.findall(".//w:t", NS)] == [node.text for node in old.findall(".//w:t", NS)]
        tables = document.findall(".//w:tbl", NS)
        assert tables
        for table in tables:
            for run in table.findall(".//w:r", NS):
                size = run.find("w:rPr/w:sz", NS)
                assert size is not None and int(size.get(W + "val")) == round(visual.typography.table_pt * 2)
                color = run.find("w:rPr/w:color", NS)
                assert color is not None and color.get(W + "val") == visual.palette.text[1:]
            for border in table.findall("w:tblPr/w:tblBorders/*", NS):
                size = visual.layout.thin_rule_pt
                if kind == "geology_docx" and border.tag.rsplit("}", 1)[-1] not in ("insideH", "insideV"):
                    size = visual.layout.strong_rule_pt
                assert int(border.get(W + "sz")) == round(size * 8)
            if kind == "geology_docx":
                old_table = old.findall(".//w:tbl", NS)[tables.index(table)]
                assert ET.tostring(table.find("w:tblGrid", NS)) == ET.tostring(old_table.find("w:tblGrid", NS))
        if kind == "generic_docx":
            caption = styles.find("w:style[@w:styleId='Caption']/w:rPr/w:sz", NS)
            assert int(caption.get(W + "val")) == round(visual.typography.caption_pt * 2)
            assert document.findall(".//w:pStyle[@w:val='Caption']", NS)
    else:
        workbook, old = load_workbook(output), load_workbook(baseline)
        try:
            assert workbook.sheetnames == old.sheetnames
            for sheet, old_sheet in zip(workbook.worksheets, old.worksheets):
                assert [(cell.value, cell.data_type) for row in sheet for cell in row] == [(cell.value, cell.data_type) for row in old_sheet for cell in row]
                assert [(key, value.width) for key, value in sheet.column_dimensions.items()] == [(key, value.width) for key, value in old_sheet.column_dimensions.items()]
                for part in (sheet.oddFooter.left, sheet.oddFooter.right):
                    assert part.size == visual.typography.footer_pt
                    assert part.color == visual.palette.text_muted[1:].upper()
                assert sheet.oddFooter.left.text == visual.brand_wordmark.replace("&", "&&")
                assert sheet.oddFooter.right.text == "&P / &N"
                if sheet is workbook.worksheets[0]:
                    assert sheet["A1"].font.sz == visual.typography.title_pt
                    for row in sheet.iter_rows(min_row=2):
                        for cell in row:
                            assert cell.font.sz == visual.typography.body_pt
                else:
                    assert sheet.page_setup.scale == 100 and sheet.page_setup.fitToWidth == 0
                    assert sheet.freeze_panes == "A2" and sheet.print_title_rows == "$1:$1"
                    assert sheet["A1"].font.sz == visual.typography.table_pt
        finally:
            workbook.close()
            old.close()
    assert (restored.current_well.cuttings, restored.current_well.stratigraphy) == before
    for key, values in source.items():
        np.testing.assert_array_equal(dataset.curves[key].values, values)
