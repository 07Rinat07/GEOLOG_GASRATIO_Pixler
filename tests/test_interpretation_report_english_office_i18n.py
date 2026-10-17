from __future__ import annotations

from dataclasses import replace
from types import SimpleNamespace
from zipfile import ZipFile

import fitz
import pytest
from openpyxl import load_workbook

from geoworkbench.data.hydrocarbon_interpretation_export import (
    export_hydrocarbon_interpretation_docx,
)
from geoworkbench.data.hydrocarbon_interpretation_export_readable import (
    _haworth_pixler_text,
    export_readable_hydrocarbon_interpretation_xlsx,
)
from geoworkbench.printing.hydrocarbon_interpretation_report import (
    export_hydrocarbon_interpretation_pdf,
)
from geoworkbench.printing.hydrocarbon_report_i18n import hydrocarbon_report_labels
from geoworkbench.services.gas_ratio_interpretation import classify_pixler_ratios
from geoworkbench.services.hydrocarbon_interpretation import (
    build_hydrocarbon_interpretation_report,
    hydrocarbon_interpretation_html,
)
from geoworkbench.services.hydrocarbon_interpretation_legacy import (
    OpusGasomerReportSection,
)
from geoworkbench.services.hydrocarbon_interpretation_report_locale import (
    localize_method_calculation,
    localize_method_name,
)
from geoworkbench.services.localization import AppLanguage
from test_readable_interpretation_export import _session_with_zero_heavy_components


@pytest.mark.parametrize("language", list(AppLanguage))
def test_pixler_office_text_uses_localized_phase_instead_of_internal_code(language):
    pixler = classify_pixler_ratios(
        c1_c2=12.0, c1_c3=24.0, c1_c4=48.0, c1_c5=96.0,
    )
    assert pixler.code == "light_oil_or_gas_condensate"
    candidate = SimpleNamespace(
        interval_wetness=12.0,
        interval_balance=8.0,
        interval_character=0.3,
        pixler_assessment=pixler,
    )
    label = _haworth_pixler_text(candidate, language)
    assert "light_oil_or_gas_condensate" not in label
    assert "C1/C2=12" in label
    assert {
        AppLanguage.RU: "переход",
        AppLanguage.KK: "мұнай",
        AppLanguage.EN: "transition: light oil or gas condensate",
    }[language] in label
    if language is AppLanguage.EN:
        assert "profile=positive" in label


@pytest.mark.parametrize("language", [AppLanguage.KK, AppLanguage.EN])
def test_actual_office_methods_match_preview_labels(tmp_path, language):
    session = _session_with_zero_heavy_components()
    dataset = session.current_dataset
    assert dataset is not None
    report = build_hydrocarbon_interpretation_report(session, threshold=3.0)
    method = report.methods[0]
    expected_name = localize_method_name(method.method, language)
    expected_calculation = localize_method_calculation(method.calculation, language)
    assert expected_name in hydrocarbon_interpretation_html(report, language)

    docx_path = tmp_path / f"interpretation-{language.value}.docx"
    export_hydrocarbon_interpretation_docx(
        report, docx_path, dataset=dataset, language=language,
    )
    with ZipFile(docx_path) as docx:
        document_xml = docx.read("word/document.xml").decode("utf-8")
    assert expected_name in document_xml
    # The calculation is always presentation text, not the raw English version
    # for Kazakh output.
    fragment = (
        "Флюид түрі палетка бойынша алдын ала бағаланады."
        if language is AppLanguage.KK
        else "The fluid label is a preliminary palette-based interpretation."
    )
    assert fragment in expected_calculation
    assert fragment in document_xml

    xlsx_path = tmp_path / f"interpretation-{language.value}.xlsx"
    export_readable_hydrocarbon_interpretation_xlsx(
        report, dataset, xlsx_path, language=language,
    )
    book = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        methods = book[hydrocarbon_report_labels(language).sheet_methods]
        assert methods.cell(2, 1).value == expected_name
        assert methods.cell(2, 4).value == expected_calculation
    finally:
        book.close()


def test_real_english_pdf_contains_english_methodology_not_russian(qapp, tmp_path):
    session = _session_with_zero_heavy_components()
    dataset = session.current_dataset
    assert dataset is not None
    report = build_hydrocarbon_interpretation_report(session, threshold=3.0)
    output = tmp_path / "english-gas-report.pdf"
    export_hydrocarbon_interpretation_pdf(
        report, output, dataset=dataset, language=AppLanguage.EN,
    )
    with fitz.open(output) as pdf:
        text = "\n".join(page.get_text() for page in pdf)
    assert "Mud-gas interpretation report" in text
    assert "Methods and availability" in text
    assert "Haworth wetness/balance/character" in text
    assert "Отчёт по интерпретации газового каротажа" not in text
    assert "Расчёт и правило интерпретации" not in text


def test_english_opus_gasomer_office_keeps_formulas_but_names_indicators(tmp_path):
    session = _session_with_zero_heavy_components()
    dataset = session.current_dataset
    assert dataset is not None
    original = build_hydrocarbon_interpretation_report(session, threshold=3.0)
    section = OpusGasomerReportSection(
        profile_id="test-opus",
        profile_version="1",
        profile_status="workbook-derived",
        calculation_mode="synchronous-rows",
        interval_source="gas show",
        working_unit="%vol",
        total_gas_lod=None,
        input_curves=(("TOTAL_GAS", "TG"), ("C1", "C1")),
        input_units=(("TOTAL_GAS", "%"), ("C1", "%")),
        formulas=(("OPUS_GM_1", "p1 * p2 / (p2 + p3)**2"),),
        class_labels=(),
        source_workbook_sha256="a" * 64,
        provenance=(),
        errata=(),
        intervals=(),
        warnings=(),
    )
    report = replace(original, report_profile="opus", opus_gasomer=section)
    docx = tmp_path / "opus-en.docx"
    export_hydrocarbon_interpretation_docx(
        report, docx, dataset=dataset, language=AppLanguage.EN,
    )
    with ZipFile(docx) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "OPUS Gasomer indicator 1" in xml
    assert "p1 * p2 / (p2 + p3)**2" in xml
    xlsx = tmp_path / "opus-en.xlsx"
    export_readable_hydrocarbon_interpretation_xlsx(
        report, dataset, xlsx, language=AppLanguage.EN,
    )
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        sheet = workbook[hydrocarbon_report_labels(AppLanguage.EN).sheet_opus]
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        assert "OPUS Gasomer indicator 1" in values
        assert "p1 * p2 / (p2 + p3)**2" in values
    finally:
        workbook.close()
