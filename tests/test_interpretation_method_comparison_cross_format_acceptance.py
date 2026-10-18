"""One immutable, operator-verified comparison across real HTML/PDF/Word/Excel.

This is a limited cross-format semantic acceptance gate. System print is
independently covered by the actual UI->temporary PDF->spooler handoff tests;
no physical device or geological-field verification is implied here.
"""

from __future__ import annotations

from dataclasses import replace
from html import escape
import json
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import fitz
import numpy as np
from openpyxl import load_workbook

from geoworkbench.data.hydrocarbon_interpretation_export_docx_polished import (
    export_polished_hydrocarbon_interpretation_docx,
)
from geoworkbench.data.hydrocarbon_interpretation_export_readable import (
    export_readable_hydrocarbon_interpretation_xlsx,
)
from geoworkbench.printing.hydrocarbon_interpretation_report import (
    export_hydrocarbon_interpretation_pdf_with_passport,
)
from geoworkbench.printing.hydrocarbon_interpretation_report_identity import (
    default_interpretation_report_identity,
)
from geoworkbench.services.hydrocarbon_interpretation import (
    build_hydrocarbon_interpretation_report,
    hydrocarbon_interpretation_html,
)
from geoworkbench.services.interpretation_method_comparison_html import (
    method_comparison_table,
)
from geoworkbench.services.localization import AppLanguage
from geoworkbench.services.report_passport import passport_sidecar_path
from test_interpretation_method_comparison_xlsx_output import _comparison
from test_readable_interpretation_export import _session_with_zero_heavy_components


_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_SHEET = {
    AppLanguage.RU: "Сравнение методов",
    AppLanguage.KK: "Әдістерді салыстыру",
    AppLanguage.EN: "Method comparison",
}


def _docx_table_cells(path: Path, headings: tuple[str, ...]) -> tuple[str, ...]:
    with ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    matches = []
    for table in root.findall(".//" + _WORD + "tbl"):
        rows = table.findall(_WORD + "tr")
        if len(rows) < 2:
            continue
        first = tuple(
            "".join(node.text or "" for node in cell.iter(_WORD + "t"))
            for cell in rows[0].findall(_WORD + "tc")
        )
        if first == headings:
            matches.append(rows[1])
    assert len(matches) == 1
    return tuple(
        "".join(node.text or "" for node in cell.iter(_WORD + "t"))
        for cell in matches[0].findall(_WORD + "tc")
    )


def _check_one_snapshot_in_all_formats(tmp_path: Path, language: AppLanguage) -> None:
    session = _session_with_zero_heavy_components()
    dataset = session.current_dataset
    assert dataset is not None
    original_depth = dataset.depth.copy()
    original_curves = {key: curve.values.copy() for key, curve in dataset.curves.items()}
    report = build_hydrocarbon_interpretation_report(session)
    original_report = replace(report)
    comparison = _comparison()
    snapshot = (comparison,)
    projection = method_comparison_table(snapshot, language, depth_unit=report.depth_unit)
    assert len(projection.rows) == 1
    identity = default_interpretation_report_identity(report, language)

    html = hydrocarbon_interpretation_html(
        report, language, method_comparisons=snapshot,
    )
    pdf = export_hydrocarbon_interpretation_pdf_with_passport(
        session, report, tmp_path / "signed.pdf",
        identity=identity, language=language, method_comparisons=snapshot,
    )
    word = export_polished_hydrocarbon_interpretation_docx(
        report, tmp_path / "client.docx", dataset=dataset,
        identity=identity, language=language, method_comparisons=snapshot,
    )
    xlsx = export_readable_hydrocarbon_interpretation_xlsx(
        report, dataset, tmp_path / "client.xlsx",
        identity=identity, language=language, method_comparisons=snapshot,
    )

    assert "method-comparison-analytics" in html
    for heading in projection.headings:
        assert escape(heading) in html
    for cell in projection.rows[0].cells:
        assert escape(cell) in html
    assert f'data-reason-code="{comparison.reason_code}"' in html

    assert _docx_table_cells(word, projection.headings) == projection.rows[0].cells

    book = load_workbook(xlsx, read_only=False, data_only=False)
    try:
        sheet = book[_SHEET[language]]
        assert tuple(sheet.cell(1, i).value for i in range(1, 6)) == projection.headings
        assert tuple(sheet.cell(2, i).value for i in range(1, 6)) == projection.rows[0].cells
        assert sheet["F2"].value == comparison.reason_code
        assert sheet["G2"].value == comparison.reason_schema_version
        assert sheet.column_dimensions["F"].hidden
        assert sheet.column_dimensions["G"].hidden
    finally:
        book.close()

    sidecar = json.loads(
        passport_sidecar_path(pdf.primary_path).read_text(encoding="utf-8")
    )
    assert dict(sidecar["render"]["options"])["method_comparison_intervals"] == "1"
    with fitz.open(pdf.primary_path) as document:
        visible_pdf = re.sub(
            r"\s+", " ", " ".join(page.get_text() for page in document)
        )
    assert projection.title in visible_pdf
    assert projection.rows[0].cells[2] in visible_pdf
    # Long table cells can wrap; verify localized explanation is present
    # without depending on PDF physical column widths.
    assert projection.rows[0].cells[3].split(" ")[0] in visible_pdf
    assert str(comparison.standard_bounds[0])[:3] in visible_pdf

    np.testing.assert_array_equal(dataset.depth, original_depth)
    for key, samples in original_curves.items():
        np.testing.assert_array_equal(dataset.curves[key].values, samples)
    assert report == original_report
    assert snapshot == (_comparison(),)


def test_one_snapshot_has_same_html_pdf_docx_xlsx_evidence_ru(tmp_path, qapp) -> None:
    _check_one_snapshot_in_all_formats(tmp_path, AppLanguage.RU)


def test_one_snapshot_has_same_html_pdf_docx_xlsx_evidence_kk(tmp_path, qapp) -> None:
    _check_one_snapshot_in_all_formats(tmp_path, AppLanguage.KK)


def test_one_snapshot_has_same_html_pdf_docx_xlsx_evidence_en(tmp_path, qapp) -> None:
    _check_one_snapshot_in_all_formats(tmp_path, AppLanguage.EN)


def test_no_operator_snapshot_stays_out_of_all_output_formats(tmp_path, qapp) -> None:
    language = AppLanguage.EN
    session = _session_with_zero_heavy_components()
    dataset = session.current_dataset
    assert dataset is not None
    report = build_hydrocarbon_interpretation_report(session)
    identity = default_interpretation_report_identity(report, language)
    expected = method_comparison_table((_comparison(),), language, depth_unit=report.depth_unit)

    html = hydrocarbon_interpretation_html(report, language)
    pdf = export_hydrocarbon_interpretation_pdf_with_passport(
        session, report, tmp_path / "plain.pdf", identity=identity, language=language,
    )
    word = export_polished_hydrocarbon_interpretation_docx(
        report, tmp_path / "plain.docx", dataset=dataset,
        identity=identity, language=language,
    )
    xlsx = export_readable_hydrocarbon_interpretation_xlsx(
        report, dataset, tmp_path / "plain.xlsx",
        identity=identity, language=language,
    )

    assert "method-comparison-analytics" not in html
    assert expected.title not in html
    with ZipFile(word) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    word_text = " ".join(node.text or "" for node in root.iter(_WORD + "t"))
    assert expected.title not in word_text
    book = load_workbook(xlsx)
    try:
        assert _SHEET[language] not in book.sheetnames
    finally:
        book.close()
    sidecar = json.loads(
        passport_sidecar_path(pdf.primary_path).read_text(encoding="utf-8")
    )
    assert "method_comparison_intervals" not in dict(sidecar["render"]["options"])
    with fitz.open(pdf.primary_path) as document:
        visible_pdf = " ".join(page.get_text() for page in document)
    assert expected.title not in visible_pdf
    assert report.dataset_id == dataset.dataset_id
