"""File adapter for WELL-02 late laboratory analysis imports."""
from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Final

from geoworkbench.data.csv_adapter import CsvImportError, CsvImportPlan, probe_csv
from geoworkbench.data.excel_adapter import ExcelImportError, probe_excel
from geoworkbench.domain.analysis_update import AnalysisField, AnalysisScalar
from geoworkbench.services.well_analysis_update import (
    AnalysisSourceSample,
    AnalysisSourceValue,
)


class LateAnalysisImportError(RuntimeError):
    """Raised when a late-analysis source cannot be converted safely."""


@dataclass(frozen=True, slots=True)
class LateAnalysisImportResult:
    source_samples: tuple[AnalysisSourceSample, ...]
    source_name: str
    source_sha256: str
    detected_fields: tuple[AnalysisField, ...]


@dataclass(frozen=True, slots=True)
class _Table:
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]


_MAX_ROWS: Final = 10_000
_NULL_TOKENS: Final = frozenset({"", "null", "none", "na", "n/a", "nan", "-", "—"})

_TOP_DEPTH_ALIASES: Final = frozenset(
    {
        "top depth",
        "depth top",
        "from",
        "from depth",
        "top",
        "кровля",
        "глубина от",
        "от",
    }
)
_BOTTOM_DEPTH_ALIASES: Final = frozenset(
    {
        "bottom depth",
        "depth bottom",
        "to",
        "to depth",
        "bottom",
        "подошва",
        "глубина до",
        "до",
    }
)
_SAMPLE_ID_ALIASES: Final = frozenset(
    {"sample id", "sample", "sample no", "sample number", "id образца", "образец"}
)

_FIELD_ALIASES: Final[dict[AnalysisField, frozenset[str]]] = {
    AnalysisField.LBA_GROUP: frozenset({"lba group", "группа лба", "лба группа"}),
    AnalysisField.LBA_TYPE_ID: frozenset(
        {"lba type id", "lba type", "тип лба", "тип люминесценции"}
    ),
    AnalysisField.LBA_INTENSITY: frozenset(
        {"lba intensity", "интенсивность лба", "интенсивность люминесценции"}
    ),
    AnalysisField.LBA_COLOR: frozenset({"lba color", "цвет лба", "цвет люминесценции"}),
    AnalysisField.LBA_DISTRIBUTION: frozenset(
        {"lba distribution", "распределение лба", "распределение люминесценции"}
    ),
    AnalysisField.LBA_CUT: frozenset({"lba cut", "вытяжка лба", "вытяжка"}),
    AnalysisField.LBA_CUT_SPEED: frozenset(
        {"lba cut speed", "скорость вытяжки", "скорость лба"}
    ),
    AnalysisField.LBA_CUT_COLOR: frozenset(
        {"lba cut color", "цвет вытяжки", "цвет вытяжки лба"}
    ),
    AnalysisField.LBA_RESIDUE_TYPE: frozenset(
        {"lba residue type", "тип остатка", "остаток лба"}
    ),
    AnalysisField.LBA_RESIDUE_COLOR: frozenset(
        {"lba residue color", "цвет остатка", "цвет остатка лба"}
    ),
    AnalysisField.LBA_ODOUR: frozenset({"lba odour", "lba odor", "запах", "запах лба"}),
    AnalysisField.LBA_STAIN: frozenset({"lba stain", "пятно", "пятно лба"}),
    AnalysisField.LBA_DESCRIPTION: frozenset(
        {"lba description", "описание лба", "описание люминесценции"}
    ),
    AnalysisField.CALCITE_PERCENT: frozenset(
        {"calcite", "calcite percent", "calcite %", "кальцит", "кальцит %", "кальцит процент"}
    ),
    AnalysisField.DOLOMITE_PERCENT: frozenset(
        {
            "dolomite",
            "dolomite percent",
            "dolomite %",
            "доломит",
            "доломит %",
            "доломит процент",
        }
    ),
    AnalysisField.ANALYSIS_INTERPRETATION: frozenset(
        {
            "analysis interpretation",
            "interpretation",
            "интерпретация анализа",
            "интерпретация",
            "заключение",
        }
    ),
}

_INTEGER_FIELDS: Final = frozenset({AnalysisField.LBA_GROUP, AnalysisField.LBA_INTENSITY})
_FLOAT_FIELDS: Final = frozenset(
    {AnalysisField.CALCITE_PERCENT, AnalysisField.DOLOMITE_PERCENT}
)


def load_late_analysis_source(path: str | Path) -> LateAnalysisImportResult:
    """Load a tabular late-analysis source into typed WELL-02 records."""

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)

    suffix = source.suffix.casefold()
    if suffix in {".csv", ".txt"}:
        table = _read_csv(source)
    elif suffix in {".xlsx", ".xlsm"}:
        table = _read_excel(source)
    else:
        raise LateAnalysisImportError(
            f"Неподдерживаемый формат поздних анализов: {source.suffix or 'без расширения'}"
        )

    mapping = _map_columns(table.headers)
    samples = tuple(
        _parse_row(row, row_number=index + 2, mapping=mapping)
        for index, row in enumerate(table.rows)
    )
    if not samples:
        raise LateAnalysisImportError("Файл поздних анализов не содержит строк данных")

    detected_fields = tuple(mapping.analysis_columns)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    return LateAnalysisImportResult(samples, source.name, digest, detected_fields)


@dataclass(frozen=True, slots=True)
class _ColumnMapping:
    top_depth: int
    bottom_depth: int
    sample_id: int | None
    analysis_columns: dict[AnalysisField, int]


def _read_csv(source: Path) -> _Table:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            probe = probe_csv(source, CsvImportPlan(encoding=encoding))
            with source.open("r", encoding=encoding, newline="") as stream:
                rows = tuple(
                    tuple(cell.strip() for cell in row)
                    for row in csv.reader(stream, delimiter=probe.delimiter)
                )
            if not rows:
                raise LateAnalysisImportError("Файл поздних анализов пуст")
            return _validated_table(tuple(rows[0]), rows[1:])
        except (CsvImportError, UnicodeDecodeError, OSError) as exc:
            last_error = exc
    raise LateAnalysisImportError(f"Не удалось прочитать {source.name}") from last_error


def _read_excel(source: Path) -> _Table:
    try:
        probe = probe_excel(source)
        if not probe.sheet_names:
            raise LateAnalysisImportError("Книга Excel не содержит листов")
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook[probe.sheet_names[0]]
            values = tuple(
                tuple("" if value is None else str(value).strip() for value in row)
                for row in sheet.iter_rows(values_only=True)
            )
        finally:
            workbook.close()
    except (ExcelImportError, ImportError, KeyError, OSError, ValueError) as exc:
        raise LateAnalysisImportError(f"Не удалось прочитать Excel: {source.name}") from exc
    if not values:
        raise LateAnalysisImportError("Книга Excel не содержит строк")
    return _validated_table(tuple(values[0]), values[1:])


def _validated_table(
    headers: tuple[str, ...],
    rows: tuple[tuple[str, ...], ...],
) -> _Table:
    if not headers or any(not header.strip() for header in headers):
        raise LateAnalysisImportError("Заголовки поздних анализов не должны быть пустыми")
    if len(rows) > _MAX_ROWS:
        raise LateAnalysisImportError(
            f"Слишком много строк анализа: {len(rows)}; максимум {_MAX_ROWS}"
        )
    width = len(headers)
    normalized_rows: list[tuple[str, ...]] = []
    for row_number, row in enumerate(rows, start=2):
        if len(row) != width:
            raise LateAnalysisImportError(
                f"Строка {row_number}: ожидалось колонок {width}, получено {len(row)}"
            )
        if any(cell.strip() for cell in row):
            normalized_rows.append(tuple(cell.strip() for cell in row))
    return _Table(tuple(header.strip() for header in headers), tuple(normalized_rows))


def _map_columns(headers: tuple[str, ...]) -> _ColumnMapping:
    normalized = tuple(_normalize_header(header) for header in headers)
    top_depth = _single_column(normalized, _TOP_DEPTH_ALIASES, "верхняя глубина")
    bottom_depth = _single_column(normalized, _BOTTOM_DEPTH_ALIASES, "нижняя глубина")
    sample_id = _optional_single_column(normalized, _SAMPLE_ID_ALIASES, "ID образца")

    analysis_columns: dict[AnalysisField, int] = {}
    for field in AnalysisField:
        aliases = {
            _normalize_header(field.value),
            *(_normalize_header(value) for value in _FIELD_ALIASES[field]),
        }
        positions = [index for index, value in enumerate(normalized) if value in aliases]
        if len(positions) > 1:
            raise LateAnalysisImportError(
                f"Поле {field.value!r} неоднозначно сопоставлено с несколькими колонками"
            )
        if positions:
            analysis_columns[field] = positions[0]
    if not analysis_columns:
        raise LateAnalysisImportError(
            "Не найдено ни одной поддерживаемой колонки позднего анализа"
        )
    return _ColumnMapping(top_depth, bottom_depth, sample_id, analysis_columns)


def _single_column(
    normalized_headers: tuple[str, ...],
    aliases: frozenset[str],
    label: str,
) -> int:
    position = _optional_single_column(normalized_headers, aliases, label)
    if position is None:
        raise LateAnalysisImportError(f"Не найдена обязательная колонка: {label}")
    return position


def _optional_single_column(
    normalized_headers: tuple[str, ...],
    aliases: frozenset[str],
    label: str,
) -> int | None:
    normalized_aliases = {_normalize_header(value) for value in aliases}
    positions = [
        index for index, value in enumerate(normalized_headers) if value in normalized_aliases
    ]
    if len(positions) > 1:
        raise LateAnalysisImportError(f"Колонка {label!r} определена неоднозначно")
    return positions[0] if positions else None


def _parse_row(
    row: tuple[str, ...],
    *,
    row_number: int,
    mapping: _ColumnMapping,
) -> AnalysisSourceSample:
    top_depth = _parse_float(row[mapping.top_depth], row_number, "верхняя глубина")
    bottom_depth = _parse_float(row[mapping.bottom_depth], row_number, "нижняя глубина")
    if bottom_depth <= top_depth:
        raise LateAnalysisImportError(
            f"Строка {row_number}: нижняя глубина должна быть больше верхней"
        )

    target_sample_id = None
    if mapping.sample_id is not None:
        candidate = row[mapping.sample_id].strip()
        if _normalize_null(candidate) is not None:
            target_sample_id = candidate

    values = tuple(
        AnalysisSourceValue(field, _parse_analysis_value(field, row[position], row_number))
        for field, position in mapping.analysis_columns.items()
    )
    try:
        return AnalysisSourceSample(top_depth, bottom_depth, values, target_sample_id)
    except ValueError as exc:
        raise LateAnalysisImportError(f"Строка {row_number}: {exc}") from exc


def _parse_analysis_value(
    field: AnalysisField,
    raw: str,
    row_number: int,
) -> AnalysisScalar | None:
    text = _normalize_null(raw)
    if text is None:
        return None
    if field in _INTEGER_FIELDS:
        number = _parse_float(text, row_number, field.value)
        if not number.is_integer():
            raise LateAnalysisImportError(
                f"Строка {row_number}, {field.value}: ожидалось целое число"
            )
        return int(number)
    if field in _FLOAT_FIELDS:
        return _parse_float(text, row_number, field.value)
    return text


def _parse_float(raw: str, row_number: int, label: str) -> float:
    text = raw.strip().replace(" ", "").replace(",", ".")
    if not text:
        raise LateAnalysisImportError(f"Строка {row_number}: {label} не заполнена")
    try:
        return float(text)
    except ValueError as exc:
        raise LateAnalysisImportError(
            f"Строка {row_number}, {label}: ожидалось число, получено {raw!r}"
        ) from exc


def _normalize_null(value: str) -> str | None:
    text = value.strip()
    return None if text.casefold() in _NULL_TOKENS else text


def _normalize_header(value: str) -> str:
    text = value.strip().casefold().replace("_", " ")
    text = re.sub(r"\[[^\]]*]|\([^)]*\)", " ", text)
    text = re.sub(r"[^0-9a-zа-яё%]+", " ", text)
    return " ".join(text.split())
