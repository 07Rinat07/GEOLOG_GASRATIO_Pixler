from __future__ import annotations

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font  # type: ignore[import-untyped]

from geoworkbench.data.spreadsheet_safety import protect_spreadsheet_row
from geoworkbench.printing.report_document_control import ReportDocumentControl
from geoworkbench.printing.report_visual_system import modern_oilfield_report_profile
from geoworkbench.services.localization import AppLanguage


def write_document_control_sheet(
    workbook: Workbook, snapshot: ReportDocumentControl, language: AppLanguage,
) -> None:
    """Write printable presentation details without changing data or audit sheets."""
    sheet = workbook.create_sheet({
        AppLanguage.RU: "Реквизиты", AppLanguage.KK: "Деректемелер", AppLanguage.EN: "Document control",
    }[language])
    visual = modern_oilfield_report_profile()
    sheet.append(protect_spreadsheet_row((visual.brand_wordmark,)))
    sheet.append(protect_spreadsheet_row((snapshot.title,)))
    sheet.append(protect_spreadsheet_row((snapshot.subtitle,)))
    for label, value in snapshot.available_rows:
        sheet.append(protect_spreadsheet_row((label, value)))
    for note in snapshot.notes:
        sheet.append(protect_spreadsheet_row((note,)))
    sheet.merge_cells("A1:B1")
    sheet.merge_cells("A2:B2")
    sheet.merge_cells("A3:B3")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 75
    for row in sheet:
        for cell in row:
            cell.font = Font(size=visual.typography.body_pt, color=visual.palette.text.lstrip("#"))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A1"].font = Font(bold=True, color=visual.palette.accent.lstrip("#"))
    sheet["A2"].font = Font(bold=True, size=visual.typography.section_pt)
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:3"
    sheet.oddFooter.left.text = visual.brand_wordmark.replace("&", "&&")
    sheet.oddFooter.right.text = "&P / &N"
